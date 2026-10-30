import json
from copy import deepcopy
from datetime import date
from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from docx import Document

from apps.accounts.models import UserRole
from apps.accounts.services import issue_api_token
from apps.evaluations.models import Evaluation
from apps.patients.models import Patient
from apps.reports.builders.tests_builder import build_validated_tests_snapshot
from apps.reports.services.report_export_service import ReportExportService
from apps.tests.base.types import TestContext
from apps.tests.models import Instrument, TestApplication
from apps.tests.registry import get_test_module
from apps.tests.services.scoring_service import TestScoringService
from apps.tests.vineland3 import Vineland3Module
from apps.tests.vineland3.calculators import compute_scores
from apps.tests.vineland3.norms import age_months, age_table, load_norms, lookup
from apps.tests.vineland3.validators import validate_input


# D.1, PDF page 8: these raw ranges correspond to v-scale 15.
PROFESSOR_SCORES = {"rec": 40, "exp": 47, "esc": 4, "pes": 30, "num": 5,
                    "cme": 19, "rin": 38, "bri": 24, "enf": 32, "cmg": 38, "cmf": 20}


def raw(form="professores", birth="2023-01-15", applied="2026-01-15"):
    table = age_table(form, "subdomains", age_months({"birth_date": birth, "reference_date": applied}))
    scores = {key: entries[0][0] for key, entries in table["scores"].items() if entries}
    return {"form": form, "respondent": "Respondente de referência", "entry_mode": "totals",
            "birth_date": birth, "reference_date": applied, "raw_scores": scores}


class Vineland3NormTests(SimpleTestCase):
    def test_registry_and_reference_case_from_d1_d3(self):
        self.assertIsInstance(get_test_module("vineland3"), Vineland3Module)
        data = compute_scores({**raw(), "raw_scores": PROFESSOR_SCORES})
        self.assertEqual({row["v_scale"] for row in data["subdomains"]}, {15})
        # D.3 PDF page 58, extensive (left) columns, not domain-level columns.
        domains = {row["scale"]: row for row in data["domains"]}
        self.assertEqual({key: row["standard_score"] for key, row in domains.items()},
                         {"COM": 99, "AVD": 98, "SOC": 100, "HMOT": 99, "ABC": 99})
        self.assertEqual(domains["ABC"]["sum_standard_scores"], 297)
        self.assertEqual(domains["ABC"]["percentile"], "47")
        self.assertEqual(data["norm_source"]["domain_pages"], [57, 58, 59])

    def test_source_coverage_and_continuity(self):
        expected = {"entrevista": (90, 30), "pais_cuidadores": (90, 30), "professores": (46, 23)}
        for form, source in load_norms()["forms"].items():
            self.assertEqual((len(source["subdomains"]), len(source["domains"])), expected[form])
            self.assertEqual(len(source["sha256"]), 64)
            for kind in ["subdomains", "domains"]:
                for previous, current in zip(source[kind], source[kind][1:]):
                    self.assertEqual(previous["max_months"] + 1, current["min_months"])
                for table in source[kind]:
                    for entries in table["scores"].values():
                        for previous, current in zip(entries, entries[1:]):
                            self.assertEqual(previous[1] + 1, current[0])

    def test_every_age_band_can_compute_minimum_and_maximum_raw_scores(self):
        for form, source in load_norms()["forms"].items():
            for table in source["subdomains"]:
                months = table["min_months"]
                applied = f"{1900 + months // 12:04d}-{1 + months % 12:02d}-01"
                for index in (0, -1):
                    scores = {key: entries[index][0 if index == 0 else 1]
                              for key, entries in table["scores"].items() if entries}
                    with self.subTest(form=form, months=months, boundary=index):
                        data = compute_scores({**raw(form, "1900-01-01", applied), "raw_scores": scores})
                        self.assertTrue(data["domains"])

    def test_infant_forms_are_distinct_and_zero_is_preserved(self):
        interview = compute_scores(raw("entrevista", "2026-01-15"))
        parents = compute_scores(raw("pais_cuidadores", "2026-01-15"))
        self.assertEqual(interview["subdomains"][0]["raw_score"], 0)
        self.assertEqual(interview["subdomains"][0]["v_scale"], 9)
        self.assertEqual(parents["subdomains"][0]["v_scale"], 7)
        self.assertNotIn("esc", {row["scale"] for row in interview["subdomains"]})

    def test_parent_coping_norms_start_at_two_not_three(self):
        table = age_table("pais_cuidadores", "subdomains", 24)
        self.assertTrue(table["scores"]["enf"])
        self.assertEqual(lookup(table["scores"]["enf"], 0)[0], 6)

    def test_age_limits_and_birthdays(self):
        self.assertEqual(age_months({"birth_date": "2020-01-31", "reference_date": "2020-02-29"}), 1)
        self.assertEqual(age_months({"birth_date": "2023-01-15", "reference_date": "2026-01-14"}), 35)
        for months in (0, 35, 228, 300):
            with self.assertRaises(ValueError):
                age_table("professores", "subdomains", months)
        self.assertTrue(age_table("professores", "subdomains", 227))
        self.assertTrue(age_table("entrevista", "subdomains", 1200))

    def test_dates_form_respondent_and_input_modes_are_validated(self):
        for changes in ({"birth_date": None}, {"reference_date": "invalid"}, {"birth_date": "2027-01-01"},
                        {"form": "invalid"}, {"respondent": " "}, {"entry_mode": "items"},
                        {"item_responses": {}}, {"raw_scores": None}):
            with self.subTest(changes=changes):
                self.assertTrue(validate_input({**raw(), **changes}))
                with self.assertRaises(ValueError):
                    compute_scores({**raw(), **changes})

    def test_strict_raw_score_validation(self):
        for value in (-1, 61, 0.5, True, "40", None):
            data = raw()
            data["raw_scores"]["rec"] = value
            self.assertTrue(validate_input(data))
        for key in ("rec", "num", "cme"):
            data = raw()
            del data["raw_scores"][key]
            self.assertTrue(validate_input(data))
        data = raw()
        data["raw_scores"]["dom"] = 1
        self.assertTrue(validate_input(data))

    def test_motor_is_optional_paired_and_age_limited(self):
        data = raw()
        del data["raw_scores"]["cmg"]
        self.assertTrue(validate_input(data))
        del data["raw_scores"]["cmf"]
        self.assertEqual(validate_input(data), [])
        self.assertNotIn("HMOT", {row["scale"] for row in compute_scores(data)["domains"]})
        data = raw(birth="2016-01-15")
        data["raw_scores"]["cmg"] = 0
        self.assertTrue(validate_input(data))

    def test_classification_is_explicitly_pending_and_text_uses_computed_scores(self):
        module = Vineland3Module()
        context = TestContext(patient_name="Paciente", evaluation_id=1, instrument_code="vineland3", raw_scores=raw())
        data = module.compute(context)
        self.assertEqual(module.classify(data), {"classification_status": "pending_manual_verification"})
        self.assertIn("normas americanas", module.interpret(context, data))
        self.assertIn("não foi emitida", module.interpret(context, data))
        self.assertEqual(module.build_report_payload(context, data)["chart_payload"]["metric"], "standard_score")
        self.assertTrue(data["requires_clinical_review"])

    def test_extreme_percentiles_are_not_converted_to_exact_percentiles(self):
        table = age_table("professores", "domains", 36)
        self.assertEqual(lookup(table["scores"]["ABC"], 60), (20, "<1"))
        self.assertEqual(lookup(table["scores"]["ABC"], 420), (140, ">99"))
        with self.assertRaises(ValueError):
            lookup(table["scores"]["ABC"], 421)


class Vineland3ApiTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="vineland-clinico", email="vineland@example.com", role=UserRole.NEUROPSYCHOLOGIST)
        self.token = issue_api_token(self.user)
        self.patient = Patient.objects.create(full_name="Paciente Vineland", birth_date=date(2023, 1, 15), created_by=self.user)
        self.evaluation = Evaluation.objects.create(patient=self.patient, examiner=self.user, start_date=date(2026, 1, 15))
        self.payload = {"evaluation_id": self.evaluation.pk, "applied_on": "2026-01-15", "form": "professores",
                        "respondent": "Professor referência", "raw_scores": deepcopy(PROFESSOR_SCORES)}

    def submit(self, **overrides):
        return self.client.post("/api/tests/vineland3/submit", data=json.dumps({**self.payload, **overrides}),
                                content_type="application/json", HTTP_AUTHORIZATION=f"Bearer {self.token}")

    def test_protocol_authentication_and_no_normative_tables_leaked(self):
        self.assertEqual(self.client.get("/api/tests/vineland3/protocol").status_code, 401)
        response = self.client.get("/api/tests/vineland3/protocol", HTTP_AUTHORIZATION=f"Bearer {self.token}")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["entry_modes"], ["totals"])
        self.assertNotIn("scores", response.json())
        scales = response.json()["forms"]["pais_cuidadores"]["scales"]
        self.assertEqual(next(row for row in scales if row["code"] == "enf")["min_months"], 24)

    def test_separate_applications_do_not_overwrite_respondents_or_forms(self):
        first = self.submit()
        self.assertEqual(first.status_code, 200, first.content)
        second = self.submit(respondent="Outro professor")
        third = self.submit(form="entrevista", raw_scores=raw("entrevista")["raw_scores"])
        self.assertEqual(third.status_code, 200, third.content)
        self.assertNotEqual(first.json()["application_id"], second.json()["application_id"])
        self.assertEqual(TestApplication.objects.filter(instrument__code="vineland3").count(), 3)
        self.assertFalse(first.json()["is_validated"])
        self.assertEqual(build_validated_tests_snapshot(self.evaluation), [])
        TestApplication.objects.filter(instrument__code="vineland3").update(is_validated=True)
        snapshot = build_validated_tests_snapshot(self.evaluation)
        self.assertEqual(len(snapshot), 3)
        self.assertIn("Outro professor", {item["respondent_name"] for item in snapshot})

    def test_edit_same_id_clears_review_and_locked_blocks_edits(self):
        app_id = self.submit().json()["application_id"]
        app = TestApplication.objects.get(pk=app_id)
        app.is_validated = True
        app.reviewed_payload = {"old": "review"}
        app.save()
        response = self.submit(application_id=app_id, respondent="Nome corrigido")
        self.assertEqual(response.status_code, 200, response.content)
        app.refresh_from_db()
        self.assertFalse(app.is_validated)
        self.assertEqual(app.reviewed_payload, {})
        self.assertIn("Nome corrigido", app.interpretation_text)
        app.status = TestApplication.Status.LOCKED
        app.save()
        self.assertEqual(self.submit(application_id=app_id).status_code, 403)

    def test_unauthorized_evaluation_and_wrong_instrument_cannot_be_updated(self):
        other = get_user_model().objects.create_user(username="vineland-other", email="other-vineland@example.com", role=UserRole.NEUROPSYCHOLOGIST)
        patient = Patient.objects.create(full_name="Outro paciente", created_by=other)
        evaluation = Evaluation.objects.create(patient=patient, examiner=other)
        self.assertEqual(self.submit(evaluation_id=evaluation.pk).status_code, 404)
        application = TestApplication.objects.create(evaluation=self.evaluation, instrument=Instrument.objects.get(code="thcp"))
        self.assertEqual(self.submit(application_id=application.pk).status_code, 404)

    def test_readonly_cannot_submit(self):
        user = get_user_model().objects.create_user(username="vineland-readonly", email="readonly-vineland@example.com", role=UserRole.READONLY)
        self.token = issue_api_token(user)
        self.assertEqual(self.submit().status_code, 403)

    def test_validation_rejects_items_bad_scores_and_future_dates(self):
        for overrides in ({"entry_mode": "items"}, {"raw_scores": {"rec": True}}, {"applied_on": "2099-01-01"}):
            self.assertEqual(self.submit(**overrides).status_code, 400)
        self.assertEqual(self.submit(birth_date="2020-01-01").status_code, 422)
        self.assertEqual(TestApplication.objects.filter(instrument__code="vineland3").count(), 0)

    def test_rescoring_uses_authoritative_dates(self):
        app = TestApplication.objects.get(pk=self.submit().json()["application_id"])
        original = deepcopy(app.computed_payload)
        app.raw_payload["birth_date"] = "1900-01-01"
        result = TestScoringService.process(app)
        self.assertTrue(result["ok"], result)
        self.assertEqual(app.computed_payload, original)

    def test_docx_uses_computed_data_and_survives_serialization(self):
        app = TestApplication.objects.get(pk=self.submit().json()["application_id"])
        app.is_validated = True
        app.interpretation_text = "Texto antigo incompatível 999"
        app.save()
        snapshot = build_validated_tests_snapshot(self.evaluation)
        self.assertEqual(len(snapshot[0]["result_rows"]), 16)
        context = {"validated_tests": snapshot}
        document = Document()
        document.add_paragraph("14. CONCLUSÃO")
        ReportExportService._insert_vineland3_results(document, context)
        ReportExportService._remove_empty_paragraphs(document)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(len(document.inline_shapes), 1)
        self.assertEqual(len(document.tables[0].rows), 17)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("escore padrão 99", text)
        self.assertNotIn("999", text)
        self.assertLess(text.index("escore padrão 99"), text.index("14. CONCLUSÃO"))
        output = BytesIO()
        document.save(output)
        restored = Document(BytesIO(output.getvalue()))
        self.assertEqual(len(restored.inline_shapes), 1)
