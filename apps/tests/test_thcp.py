import json
import math
from datetime import date
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.test import SimpleTestCase, TestCase
from docx import Document

from apps.accounts.models import UserRole
from apps.accounts.services import issue_api_token
from apps.evaluations.models import Evaluation
from apps.patients.models import Patient
from apps.reports.builders.tests_builder import build_result_rows, build_validated_tests_snapshot
from apps.reports.services.report_export_service import ReportExportService
from apps.tests.base.types import TestContext
from apps.tests.models import Instrument, TestApplication
from apps.tests.registry import get_test_module
from apps.tests.services.pdf_export_service import TestPdfExportService
from apps.tests.services.report_payload_service import TestReportPayloadService
from apps.tests.services.scoring_service import TestScoringService
from apps.tests.thcp import THCPModule
from apps.tests.thcp.classifiers import classify_z_score
from apps.tests.thcp.config import SCALES
from apps.tests.thcp.norms import SUBSCALE_NORMS, TOTAL_T_SCORES, lookup_subscale, lookup_total
from apps.tests.thcp.pdf_service import THCPPdfService


RAW = {"hpm": 22, "linguagem": 11, "pq": 9, "memoria": 7, "atencao_acertos": 26, "atencao_erros": 1, "norm_type": "idade"}


def context(age=6, **overrides):
    return TestContext(patient_name="Paciente THCP", evaluation_id=1, instrument_code="thcp", patient_age=age, raw_scores={**RAW, **overrides})


def corrected(ctx):
    module = THCPModule()
    computed = module.compute(ctx)
    return {**computed, **module.classify(computed)}


class THCPModuleTests(SimpleTestCase):
    def test_registered_and_computes_all_scales(self):
        module = get_test_module("thcp")
        self.assertIsInstance(module, THCPModule)
        self.assertEqual(module.validate(context()), [])
        data = corrected(context())
        self.assertEqual(data["scores"]["atencao"], 25)
        self.assertEqual(data["scores"]["total"], 74)
        self.assertEqual(data["summary"]["t_score"], 50)
        self.assertEqual(data["summary"]["percentile"], 50)
        self.assertEqual(data["summary"]["classification"], "Média")
        self.assertEqual(len(data["results"]), 6)

    def test_attention_clamps_negative_score_without_changing_other_scales(self):
        data = corrected(context(atencao_acertos=0, atencao_erros=62))
        self.assertEqual(data["scores"]["atencao"], 0)
        self.assertEqual(data["scores"]["total"], 49)

    def test_validates_missing_fractional_negative_excessive_and_boolean_scores(self):
        module = THCPModule()
        for field, maximum in [(key, max_score) for key, (_, max_score) in SCALES.items() if key not in {"total", "atencao"}] + [("atencao_acertos", 28), ("atencao_erros", 62)]:
            for value in (-1, maximum + 1, 0.5, True, None, "2"):
                with self.subTest(field=field, value=value):
                    self.assertTrue(module.validate(context(**{field: value})))
        ctx = context()
        ctx.raw_scores.pop("hpm")
        self.assertTrue(module.validate(ctx))

    def test_age_and_norm_validation(self):
        for age in (0, 3, 8, 12):
            self.assertTrue(THCPModule().validate(context(age=age)))
        for age in (4, 5, 6, 7):
            self.assertEqual(THCPModule().validate(context(age=age)), [])
        self.assertTrue(THCPModule().validate(context(norm_type="invalid")))

    def test_general_sample_uses_mean_sd_and_has_no_manual_total_classification(self):
        data = corrected(context(norm_type="geral"))
        z = (74 - 64.1) / 15.6
        self.assertAlmostEqual(data["summary"]["t_score"], 50 + 10 * z, places=2)
        self.assertAlmostEqual(data["summary"]["percentile"], 50 * (1 + math.erf(z / math.sqrt(2))), places=2)
        self.assertIsNone(data["summary"]["manual_classification"])
        self.assertEqual(data["norm_label"], "Amostra Geral")

    def test_manual_and_z_classifications_are_distinct(self):
        data = corrected(context())
        hpm = data["results"][0]
        self.assertEqual(hpm["manual_classification"], "Média")
        self.assertEqual(hpm["manual_percentile_band"], "50")
        self.assertEqual(hpm["percentile"], 54.55)
        high_hpm = corrected(context(hpm=25))["results"][0]
        self.assertEqual(high_hpm["manual_classification"], "Superior")
        self.assertEqual(high_hpm["classification"], "Média Superior")

    def test_z_boundaries_match_workbook(self):
        for z, expected in [(2, "Muito Superior"), (1.333, "Superior"), (0.666, "Média Superior"), (-0.666, "Média"), (-1.333, "Média Inferior"), (-2, "Limítrofe"), (-2.001, "Deficitário"), (None, None)]:
            self.assertEqual(classify_z_score(z), expected)

    def test_censored_total_preserves_interval_without_fabricating_metrics(self):
        for age, expected in [(4, "< 27"), (5, "< 24"), (6, "< 25"), (7, "< 28")]:
            ctx = context(age=age, hpm=0, linguagem=0, pq=0, memoria=0, atencao_acertos=0, atencao_erros=0)
            data = corrected(ctx)
            total = data["summary"]
            self.assertEqual(total["t_score"], expected)
            self.assertEqual(total["manual_classification"], "Muito Baixo")
            for field in ("z_score", "weighted_score", "percentile", "classification"):
                self.assertIsNone(total[field])
            self.assertIn(expected, THCPModule().interpret(ctx, data))
            self.assertTrue(data["warnings"])

    def test_maximum_total_is_91_and_preserves_upper_interval(self):
        for age, expected in [(4, "> 73"), (5, "> 76"), (6, "> 75"), (7, "> 73")]:
            data = corrected(context(age=age, hpm=30, linguagem=12, pq=11, memoria=10, atencao_acertos=28, atencao_erros=0))
            self.assertEqual(data["scores"]["total"], 91)
            self.assertEqual(data["summary"]["t_score"], expected)
            self.assertIsNone(data["summary"]["percentile"])

    def test_source_anomaly_is_preserved_and_flagged(self):
        data = corrected(context(hpm=20, linguagem=10, pq=9, memoria=10, atencao_acertos=0, atencao_erros=0))
        self.assertEqual(data["summary"]["raw_score"], 49)
        self.assertEqual(data["summary"]["t_score"], "< 28")
        self.assertIn("conferir o manual", data["warnings"][-1])

    def test_norms_have_complete_coverage(self):
        for group, values in TOTAL_T_SCORES.items():
            self.assertEqual(len(values), 92, group)
            for score in range(92):
                self.assertTrue(lookup_total(group, score)["manual_classification"])
        for group in SUBSCALE_NORMS:
            for code, (_, maximum) in SCALES.items():
                if code == "total":
                    continue
                for score in range(maximum + 1):
                    self.assertTrue(lookup_subscale(group, code, score)["manual_classification"])

    def test_all_norms_match_source_workbook_cells(self):
        path = Path(__file__).resolve().parents[2] / "aux" / "CORRECAO.xlsm"
        if not path.exists():
            self.skipTest("Planilha de referência não disponível; normas embarcadas testadas separadamente.")
        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        with ZipFile(path) as archive:
            strings = ["".join(si.itertext()) for si in ET.fromstring(archive.read("xl/sharedStrings.xml"))]
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            sheet = next(s for s in workbook.find("s:sheets", ns) if s.get("name") == "THCP-Normas")
            rel_id = sheet.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
            rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
            target = next(r.get("Target") for r in rels if r.get("Id") == rel_id)
            xml_path = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
            cells = {}
            for cell in ET.fromstring(archive.read(xml_path)).findall(".//s:c", ns):
                value = cell.find("s:v", ns)
                if value is not None:
                    cells[cell.get("r")] = strings[int(value.text)] if cell.get("t") == "s" else float(value.text)
        def col_name(number):
            result = ""
            while number:
                number, rem = divmod(number - 1, 26)
                result = chr(65 + rem) + result
            return result
        for group, start in [("geral", 3), ("4", 29), ("5", 55), ("6", 81), ("7", 107)]:
            for offset, code in enumerate(list(SCALES)[:5]):
                col = col_name(start + offset)
                for row in range(11, 103):
                    if f"{col}{row}" not in cells:
                        continue
                    score = int(cells[f"{col}{row}"])
                    norm = lookup_subscale(group, code, score)
                    self.assertEqual(norm["manual_percentile_band"], str(cells[f"A{row}"]).removesuffix(".0"))
                    self.assertEqual(norm["manual_classification"], cells[f"B{row}"])
                    self.assertEqual(norm["mean"], cells[f"{col}103"])
                    self.assertEqual(norm["sd"], cells[f"{col}104"])
        for group, col in [("4", "C"), ("5", "D"), ("6", "E"), ("7", "F")]:
            for row in range(111, 358):
                if f"{col}{row}" in cells:
                    norm = lookup_total(group, int(cells[f"{col}{row}"]))
                    self.assertEqual(norm["t_score"], cells[f"B{row}"])
                    self.assertEqual(norm["manual_classification"], cells[f"A{row}"])

    def test_report_pdf_and_docx_use_same_results_and_escape_patient_name(self):
        data = corrected(context())
        app = SimpleNamespace(computed_payload=data, classified_payload=data, applied_on=date(2026, 10, 5), evaluation=SimpleNamespace(patient=SimpleNamespace(full_name="Paciente <script>")))
        html = THCPPdfService.render_html(app)
        self.assertIn("Paciente &lt;script&gt;", html)
        self.assertIn("22/30", html)
        self.assertIn("54.55%", html)
        self.assertIn("thcp", TestPdfExportService.EXPORTERS)


class THCPApiTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user(username="thcp-clinico", email="thcp-clinico@example.com", role=UserRole.NEUROPSYCHOLOGIST, password="test-secret")
        self.token = issue_api_token(self.user)
        self.patient = Patient.objects.create(full_name="Paciente THCP", birth_date=date(2020, 10, 5), sex="F")
        self.evaluation = Evaluation.objects.create(patient=self.patient, examiner=self.user, start_date=date(2026, 10, 5))
        self.instrument = Instrument.objects.get(code="thcp")

    def submit(self, **overrides):
        return self.client.post("/api/tests/thcp/submit", data=json.dumps({"evaluation_id": self.evaluation.pk, **RAW, **overrides}), content_type="application/json", HTTP_AUTHORIZATION=f"Bearer {self.token}")

    def test_submit_result_and_snapshot(self):
        response = self.submit()
        self.assertEqual(response.status_code, 200, response.content)
        app_id = response.json()["application_id"]
        app = TestApplication.objects.get(pk=app_id)
        self.assertEqual(app.status, "reviewed")
        self.assertEqual(app.computed_payload["age"], 6)
        result = self.client.get(f"/api/tests/thcp/result/{app_id}", HTTP_AUTHORIZATION=f"Bearer {self.token}")
        self.assertEqual(result.status_code, 200)
        self.assertEqual(len(result.json()["report_payload"]["results"]), 6)
        snapshot = build_validated_tests_snapshot(self.evaluation)
        self.assertEqual(snapshot[0]["instrument_code"], "thcp")
        self.assertEqual(len(snapshot[0]["result_rows"]), 6)
        self.assertIn("74", snapshot[0]["summary"])

    def test_edition_updates_same_application_and_replaces_interpretation(self):
        app_id = self.submit().json()["application_id"]
        response = self.submit(application_id=app_id, hpm=23)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(TestApplication.objects.count(), 1)
        app = TestApplication.objects.get(pk=app_id)
        self.assertEqual(app.computed_payload["scores"]["total"], 75)
        self.assertIn("75 de 91", app.interpretation_text)

    def test_generic_rescoring_uses_application_age_and_preserves_report_zero(self):
        app_id = self.submit().json()["application_id"]
        app = TestApplication.objects.get(pk=app_id)
        original = app.computed_payload
        result = TestScoringService.process(app)
        self.assertTrue(result["ok"], result)
        self.assertEqual(app.computed_payload, original)
        app.raw_payload = {**RAW, "hpm": 0}
        app.save()
        TestScoringService.process(app)
        self.assertEqual(TestReportPayloadService.build_for_application(app)["results"][0]["raw_score"], 0)

    def test_age_is_from_application_date_not_today_or_evaluation_date(self):
        response = self.submit(applied_on="2026-10-04")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["computed_payload"]["age"], 5)
        for value in ("2024-10-04", "2028-10-05"):
            self.assertEqual(self.submit(applied_on=value).status_code, 400)

    def test_missing_birth_date_and_out_of_range_age_are_rejected(self):
        self.patient.birth_date = None
        self.patient.save()
        self.assertEqual(self.submit().status_code, 400)
        self.patient.birth_date = date(2010, 10, 5)
        self.patient.save()
        self.assertEqual(self.submit().status_code, 400)
        self.assertEqual(TestApplication.objects.count(), 0)

    def test_invalid_scores_are_rejected_without_creating_application(self):
        for changes in ({"hpm": -1}, {"hpm": 31}, {"hpm": 2.5}, {"memoria": True}, {"norm_type": "invalid"}):
            self.assertEqual(self.submit(**changes).status_code, 422)
        self.assertEqual(TestApplication.objects.count(), 0)

    def test_authentication_permissions_and_locked_application(self):
        self.assertEqual(self.client.post("/api/tests/thcp/submit", data=json.dumps({"evaluation_id": self.evaluation.pk, **RAW}), content_type="application/json").status_code, 401)
        app_id = self.submit().json()["application_id"]
        app = TestApplication.objects.get(pk=app_id)
        app.status = TestApplication.Status.LOCKED
        app.save()
        self.assertEqual(self.submit(application_id=app_id, hpm=23).status_code, 403)
        readonly = get_user_model().objects.create_user(username="thcp-readonly", email="thcp-readonly@example.com", role=UserRole.READONLY)
        self.token = issue_api_token(readonly)
        self.assertEqual(self.submit().status_code, 403)

    def test_cannot_update_another_evaluation_or_instrument(self):
        other_evaluation = Evaluation.objects.create(patient=self.patient, examiner=self.user, start_date=date(2026, 10, 5))
        other = TestApplication.objects.create(evaluation=other_evaluation, instrument=self.instrument)
        self.assertEqual(self.submit(application_id=other.pk).status_code, 404)
        other.instrument = Instrument.objects.get(code="bpa2")
        other.evaluation = self.evaluation
        other.save()
        self.assertEqual(self.submit(application_id=other.pk).status_code, 404)

    def test_catalog_exposes_age_and_name(self):
        response = self.client.get("/api/tests/instruments", HTTP_AUTHORIZATION=f"Bearer {self.token}")
        row = next(row for row in response.json() if row["code"] == "thcp")
        self.assertEqual((row["min_age"], row["max_age"]), (4, 7))
        self.assertIn("Pré-Alfabetização", row["name"])

    def test_docx_contains_table_chart_and_consistent_interpretation_before_conclusion(self):
        self.submit()
        snapshot = build_validated_tests_snapshot(self.evaluation)
        document = Document()
        document.add_paragraph("ANÁLISE QUALITATIVA")
        document.add_paragraph("14. CONCLUSÃO")
        ctx = {"validated_tests": snapshot, "patient": {"full_name": self.patient.full_name}}
        ReportExportService._insert_thcp_results(document, ctx)
        ReportExportService._sanitize_generated_document(document, None, ctx)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(len(document.tables[0].rows), 7)
        self.assertEqual(document.tables[0].cell(6, 1).text, "74/91")
        self.assertEqual(len(document.inline_shapes), 1)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertLess(text.index("74 de 91"), text.index("14. CONCLUSÃO"))
        self.assertEqual(len(build_result_rows("thcp", snapshot[0]["classified_payload"])), 6)
        output = BytesIO()
        document.save(output)
        self.assertEqual(len(Document(BytesIO(output.getvalue())).inline_shapes), 1)

    def test_docx_without_conclusion_inserts_thcp_before_references(self):
        self.submit()
        ctx = {"validated_tests": build_validated_tests_snapshot(self.evaluation), "patient": {"full_name": self.patient.full_name}}
        document = Document()
        document.add_paragraph("ANÁLISE QUALITATIVA")
        document.add_paragraph("17. REFERÊNCIAS BIBLIOGRÁFICAS")
        ReportExportService._insert_thcp_results(document, ctx)
        ReportExportService._sanitize_generated_document(document, None, ctx)
        self.assertEqual(len(document.tables), 1)
        self.assertEqual(len(document.inline_shapes), 1)
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertLess(text.index("74 de 91"), text.index("17. REFERÊNCIAS"))

    def test_pdf_export_dispatches_to_thcp_service(self):
        app_id = self.submit().json()["application_id"]
        app = TestApplication.objects.get(pk=app_id)
        with patch("apps.tests.thcp.pdf_service.generate_pdf_from_html", return_value=b"%PDF-test") as renderer:
            self.assertEqual(TestPdfExportService.build_pdf_bytes(app), b"%PDF-test")
        self.assertIn("Escore Total", renderer.call_args.args[0])
