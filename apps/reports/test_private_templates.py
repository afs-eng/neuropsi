from copy import deepcopy
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase
from django.utils.datastructures import MultiValueDict
from docx import Document
from ninja.testing import TestClient
from apps.accounts.models import hash_api_token
from unittest.mock import patch

from apps.reports.api.template_endpoints import router
from apps.reports.models import ReportTemplate
from apps.reports.services.professional_identity_service import ProfessionalIdentityService
from apps.reports.services.report_template_service import ReportTemplateService
from apps.reports.services.specialist_planning_service import SpecialistPlanningService


def docx_bytes():
    buffer = BytesIO()
    Document().save(buffer)
    return buffer.getvalue()


class PrivateTemplateTests(TestCase):
    def setUp(self):
        self.owner = get_user_model().objects.create_user(username="owner", email="owner@example.com", role="neuropsychologist", api_token_hash=hash_api_token("owner-secret"))
        self.other = get_user_model().objects.create_user(username="other", email="other@example.com", role="admin", api_token_hash=hash_api_token("other-secret"))
        self.owner_headers = {"Authorization": "Bearer owner-secret"}
        self.other_headers = {"Authorization": "Bearer other-secret"}
        self.template = ReportTemplate.objects.create(owner=self.owner, name="Infantil", instrument_code="wisc4")
        self.client = TestClient(router)

    def add_version(self, owner=None):
        return ReportTemplateService.add_version(
            owner=owner or self.owner, template_id=self.template.pk,
            filename="modelo.docx", data=docx_bytes(),
        )

    def test_versions_preserve_original_bytes_and_are_not_exportable(self):
        data = docx_bytes()
        with patch("apps.reports.test_private_templates.docx_bytes", return_value=data):
            first = self.add_version()
            second = self.add_version()
        first.refresh_from_db()
        self.assertEqual((first.version_number, second.version_number), (1, 2))
        self.assertEqual(bytes(first.document), bytes(second.document))
        self.assertEqual(bytes(first.document), data)
        self.assertEqual(first.validation_payload["status"], "draft")
        self.assertFalse(first.validation_payload["export_enabled"])

    def test_admin_cannot_append_to_another_users_template(self):
        with self.assertRaises(ReportTemplate.DoesNotExist):
            self.add_version(owner=self.other)
        self.assertFalse(self.template.versions.exists())

    def test_list_and_download_are_owner_scoped_even_for_admin(self):
        version = self.add_version()
        self.assertEqual(self.client.get("/", headers=self.other_headers).json(), [])
        self.assertEqual(len(self.client.get("/", headers=self.owner_headers).json()), 1)
        url = f"/{self.template.pk}/versions/{version.pk}/download"
        self.assertEqual(self.client.get(url, headers=self.other_headers).status_code, 404)
        response = self.client.get(url, headers=self.owner_headers)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.content, bytes(version.document))
        self.assertEqual(response["Cache-Control"], "private, no-store")

    def test_create_owner_is_not_user_controlled(self):
        response = self.client.post("/", json={"name": "Adulto", "owner_id": self.other.pk}, headers=self.owner_headers)
        self.assertEqual(response.status_code, 201)
        self.assertEqual(ReportTemplate.objects.get(pk=response.json()["id"]).owner, self.owner)

    def test_authentication_required(self):
        self.assertEqual(self.client.get("/").status_code, 401)

    def test_upload_is_private_and_cannot_enable_export(self):
        url = f"/{self.template.pk}/versions"
        def files():
            return MultiValueDict({"file": [SimpleUploadedFile("modelo.docx", docx_bytes())]})
        denied = self.client.post(url, FILES=files(), headers=self.other_headers)
        self.assertEqual(denied.status_code, 404)
        response = self.client.post(url, FILES=files(), headers=self.owner_headers)
        self.assertEqual(response.status_code, 201)
        self.assertFalse(response.json()["validation"]["export_enabled"])

    def test_blank_name_rejected(self):
        response = self.client.post("/", json={"name": "  "}, headers=self.owner_headers)
        self.assertEqual(response.status_code, 400)


class DocxInspectionTests(SimpleTestCase):
    def archive(self, extra_name, content):
        buffer = BytesIO()
        with ZipFile(BytesIO(docx_bytes())) as original, ZipFile(buffer, "w") as output:
            for name in original.namelist():
                output.writestr(name, original.read(name))
            output.writestr(extra_name, content)
        return buffer.getvalue()

    def test_valid_docx_is_draft(self):
        self.assertEqual(ReportTemplateService.inspect_docx(docx_bytes())["status"], "draft")

    def test_rejects_invalid_oversized_and_unsafe_archives(self):
        for data in (
            b"not a docx", b"x" * (ReportTemplateService.MAX_UPLOAD_BYTES + 1),
            self.archive("../escape", b"x"),
            self.archive("word/vbaProject.bin", b"x"),
            self.archive("word/activeX/activeX1.bin", b"x"),
            self.archive("word/extra.xml", b'<!DOCTYPE x [<!ENTITY e "test">]><x>&e;</x>'),
            self.archive("word/_rels/extra.rels", b'<Relationships><Relationship TargetMode="External" Target="https://example.com"/></Relationships>'),
        ):
            with self.subTest(size=len(data)), self.assertRaises(ValueError):
                ReportTemplateService.inspect_docx(data)


class SpecialistPlanningTests(SimpleTestCase):
    def setUp(self):
        self.context = {"patient": {"id": 1}, "evaluation": {"id": 2}, "validated_tests": [
            {"id": 10, "instrument_code": "wisc4", "is_validated": True, "computed_payload": {"qi": 100}},
            {"id": 11, "instrument_code": "scared", "is_validated": True},
            {"id": 12, "instrument_code": "scared", "is_validated": True},
            {"id": 13, "instrument_code": "fdt", "is_validated": False},
        ]}

    def test_one_plan_per_validated_application_without_inventing_tests(self):
        plan = SpecialistPlanningService.build_plan(self.context)
        self.assertEqual([item["application_id"] for item in plan["applications"]], [10, 11, 12])
        self.assertNotEqual(plan["applications"][1]["source_sha256"], plan["applications"][2]["source_sha256"])
        self.assertFalse(plan["execution_enabled"])

    def test_result_or_patient_change_invalidates_fingerprint_without_mutation(self):
        original = deepcopy(self.context)
        before = SpecialistPlanningService.build_plan(self.context)
        self.assertEqual(original, self.context)
        self.context["validated_tests"][0]["computed_payload"]["qi"] = 110
        after = SpecialistPlanningService.build_plan(self.context)
        self.assertNotEqual(before["applications"][0]["source_sha256"], after["applications"][0]["source_sha256"])
        self.context["patient"]["id"] = 3
        other = SpecialistPlanningService.build_plan(self.context)
        self.assertNotEqual(after["applications"][0]["source_sha256"], other["applications"][0]["source_sha256"])

    def test_unknown_instrument_is_not_silently_assigned_a_generic_specialist(self):
        plan = SpecialistPlanningService.build_plan({"validated_tests": [{"instrument_code": "unknown", "is_validated": True}]})
        self.assertEqual(plan["applications"][0]["status"], "unsupported")

    def test_identity_comes_from_report_author_not_logged_in_user(self):
        report = SimpleNamespace(author=SimpleNamespace(full_name="Profissional Exemplo", crp="01/12345"))
        self.assertEqual(ProfessionalIdentityService.author_label(report), "Profissional Exemplo (CRP 01/12345)")
        self.assertEqual(ProfessionalIdentityService.author_label(SimpleNamespace()), "Profissional não informado")
        report.evaluation = SimpleNamespace(examiner=SimpleNamespace(full_name="Responsável Clínico", crp="02/54321"))
        self.assertEqual(ProfessionalIdentityService.author_label(report), "Responsável Clínico (CRP 02/54321)")

    def test_wisc_and_wais_exports_use_actual_author_in_identification_and_signature(self):
        from apps.reports.tests import ReportExportChartSanitizationTests
        from apps.reports.services.report_export_service import ReportExportService

        for code in ("wisc4", "wais3"):
            with self.subTest(code=code):
                context = ReportExportChartSanitizationTests._export_context(self, code, 17)
                report = SimpleNamespace(
                    author=SimpleNamespace(full_name="Ana Carolina Exemplo", crp="01/12345"),
                    sections=SimpleNamespace(all=lambda: []), patient=None,
                )
                with patch("apps.reports.services.report_export_service.ReportContextService.sync_report_context", return_value=context):
                    data = ReportExportService.generate_docx_bytes(report)
                text = "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)
                self.assertEqual(text.count("Ana Carolina Exemplo (CRP 01/12345)"), 2)
                self.assertNotIn("Jacqueline", text)
