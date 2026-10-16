"""Preenche o modelo WISC-IV sem reconstruir sua apresentação."""

from copy import deepcopy
from datetime import date
import re
import unicodedata

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell, _Row
from docx.text.paragraph import Paragraph
from lxml import etree


class WISC4DocxBuilder:
    # Intervalos do modelo versionado: mantemos os objetos originais, inclusive
    # tabelas, campos SEQ, desenhos, estilos e parágrafos em branco.
    OPTIONAL_SECTIONS = {
        "bpa2": (74, 85),
        "ravlt": (85, 100),
        "fdt": (100, 116),
        "etdah_pais": (116, 131),
        "scared": (131, 150),
        "epq_j": (150, 162),
        "srs2": (162, 175),
    }

    def __init__(self, service, report, context, sections):
        self.service = service
        self.report = report
        self.context = context
        self.sections = sections
        self.document = Document(service.WISC4_TEMPLATE_PATH)
        self.paragraphs = list(self.document.paragraphs)
        self.tables = list(self.document.tables)
        self._validate_model()

    def _validate_model(self):
        expected = {3: "IDENTIFICAÇÃO", 34: "EFICIÊNCIA COGNITIVA", 48: "Função Executiva", 175: "Conclusão", 213: "Referencia Bibliográfica"}
        if len(self.paragraphs) != 225 or len(self.tables) != 12 or any(
            self.paragraphs[index].text.strip() != heading for index, heading in expected.items()
        ):
            raise ValueError("Modelo WISC-IV alterado: revise os vínculos antes de exportar, para não manter dados de exemplo.")

    @staticmethod
    def _normal(text):
        normalized = "".join(char for char in unicodedata.normalize("NFKD", str(text)).casefold() if char.isalnum())
        return {
            "sequenciadenumeroseletras": "seqnumeletras",
            "seqnumeroseletras": "seqnumeletras",
        }.get(normalized, normalized)

    def _text(self, paragraph, text):
        """Troca texto, conservando pPr e a tipografia original dos runs."""
        runs = [run for run in paragraph.runs if run.text.strip()]
        normal = next((run for run in runs if not run.bold), runs[0] if runs else None)
        properties = deepcopy(normal._r.rPr) if normal is not None and normal._r.rPr is not None else None
        paragraph.clear()
        text = re.sub(r"^(Interpretação(?: e Observações Clínicas)?|Capacidade Cognitiva Global):", r"**\1:**", str(text or ""))
        for piece in re.split(r"(\*\*.*?\*\*)", text):
            if not piece:
                continue
            bold = piece.startswith("**") and piece.endswith("**")
            run = paragraph.add_run(piece[2:-2] if bold else piece)
            if properties is not None:
                run._r.insert(0, deepcopy(properties))
            if bold:
                run.bold = True

    def _replace(self, start, stop, text, prototype=None):
        """Substitui só parágrafos; não toca nas tabelas/desenhos entre eles."""
        original = self.paragraphs[start:stop]
        anchor = original[0]._p
        if anchor.getparent() is None:
            return
        source = self.paragraphs[prototype if prototype is not None else start]
        lines = [line.strip() for line in str(text or "").splitlines() if line.strip()]
        for line in lines:
            clone = deepcopy(source._p)
            paragraph = Paragraph(clone, self.document._body)
            self._text(paragraph, line)
            anchor.addprevious(clone)
        for paragraph in original:
            element = paragraph._p
            if element.getparent() is not None:
                element.getparent().remove(element)

    def _remove_section(self, start, stop):
        first = self.paragraphs[start]._p
        end = self.paragraphs[stop]._p
        element = first
        while element is not None and element is not end:
            following = element.getnext()
            element.getparent().remove(element)
            element = following

    def _fill_table(self, index, rows, headers=1, key=None):
        table = self.tables[index]
        if key:
            rows = self.service._with_table_title(rows, key) if rows else None
        if not rows or len(rows) <= headers:
            if table._tbl.getparent() is not None:
                table._tbl.getparent().remove(table._tbl)
            return
        if any(len(row) != len(table.columns) for row in rows):
            raise ValueError("Resultados incompatíveis com as colunas do modelo WISC-IV.")
        templates = list(table.rows[headers:])
        if not templates:
            raise ValueError("Modelo WISC-IV sem linhas de dados.")
        by_label = {self._normal(row.cells[0].text): row for row in templates}
        for row in list(table.rows[headers:]):
            table._tbl.remove(row._tr)
        for position, values in enumerate(rows[headers:]):
            matching = by_label.get(self._normal(values[0]))
            source = matching if matching is not None else templates[min(position, len(templates) - 1)]
            values = list(values)
            if matching is not None:
                # Nomes abreviados do modelo são intencionais para caber na coluna.
                # Somente rótulos equivalentes são mantidos; resultados são novos.
                values[0] = source.cells[0].text
            row = _Row(deepcopy(source._tr), table)
            column = 0
            for tc in row._tr.tc_lst:
                cell = _Cell(tc, table)
                span = cell._tc.grid_span
                text = "\n".join(str(value) for value in values[column:column + span] if str(value))
                prototype = deepcopy(cell.paragraphs[0]._p)
                for child in list(cell._tc):
                    if child.tag != qn("w:tcPr"):
                        cell._tc.remove(child)
                cell._tc.append(prototype)
                self._text(Paragraph(prototype, cell), text)
                column += span
            table._tbl.append(row._tr)

    def _interpretation(self, code, section=None, fallback=None):
        return self.service._resolve_interpretation_text(
            self.sections.get(section or code), self.sections.get(fallback) if fallback else None,
            self.service._find_test(self.context, code), self.context,
        )

    def build(self):
        service = self.service
        patient = self.context.get("patient") or {}
        evaluation = self.context.get("evaluation") or {}
        for index, label, value in (
            (5, "Autora", service.FIXED_AUTHOR),
            (6, "Interessado", getattr(self.report, "interested_party", "") or service.FIXED_INTERESTED_PARTY),
            (7, "Finalidade", service.FIXED_PURPOSE),
            (9, "Nome", patient.get("full_name") or "Não informado"),
            (10, "Sexo", service._format_sex_display(patient.get("sex"))),
            (11, "Data de nascimento", service._format_date_display(patient.get("birth_date"))),
            (12, "Idade", service._age_text(self.context)),
            (13, "Filiação", service._filiation_text(patient)),
            (14, "Escolaridade", service._schooling_text(patient)),
            (15, "Colégio", patient.get("school_name") or "Não informado"),
        ):
            self._text(self.paragraphs[index], f"{label}: {value}")
        if patient.get("sex") != "F":
            self._text(self.paragraphs[8], "1.2. Identificação do paciente:")
            self._text(self.paragraphs[40], "Desempenho do paciente no WISC-IV")
        self._replace(18, 19, self.sections.get("descricao_demanda") or evaluation.get("referral_reason") or "Sem conteúdo disponível para esta seção.")
        procedures = service._wais3_procedures_intro_text(self.report, self.sections)
        self._replace(20, 21, procedures)
        self._replace(21, 30, "\n".join(f"**{item['name']}**: {item['description']}" for item in service._adolescent_instruments(self.context)), prototype=22)
        self._replace(32, 34, self.sections.get("historia_pessoal") or "Sem conteúdo disponível para esta seção.")
        test = service._find_test(self.context, "wisc4")
        self._replace(35, 36, service._wisc_global_intro_text(test, self.context))
        indices = service._wisc_global_bullet_parts(test)
        self._replace(36, 40, "\n".join(f"**{lead}**{tail}" for lead, tail in indices))
        self._replace(43, 46, service._wisc_model_interpretation_text(test, self.context))
        for index, domain, start, stop in (
            (0, "funcoes_executivas", 50, 54), (1, "linguagem", 56, 61),
            (2, "gnosias_praxias", 63, 67), (3, "memoria_aprendizagem", 69, 74),
        ):
            rows = service._wisc_memory_rows(test, self.context) if index == 3 else service._wisc_rows(test, self.context, domain)
            self._fill_table(index, rows)
            self._replace(start, stop, service._wisc_section_text(self.sections, domain, self.context))
        self._optional_results()
        self._additional_results()
        self._replace(176, 185, self.sections.get("conclusao") or self.sections.get("hipotese_diagnostica") or "Sem conteúdo disponível para esta seção.")
        self._replace(187, 194, self.sections.get("sugestoes_conduta") or self.sections.get("conduta_encaminhamentos") or "Sem conteúdo disponível para esta seção.")
        self._text(self.paragraphs[200], f"Goiânia, {service._format_date_display(str(evaluation.get('end_date') or date.today().isoformat()))}.")
        self._replace(214, 225, "\n".join(service._references_list(self.context)))
        self._bind_charts()
        self._renumber_captions()
        return self.document

    def _optional_results(self):
        service = self.service
        for code, (start, stop) in self.OPTIONAL_SECTIONS.items():
            if not service._find_test(self.context, code):
                self._remove_section(start, stop)
        test = service._find_test(self.context, "bpa2")
        if test:
            self._fill_table(4, service._bpa_rows(test, self.context))
            self._replace(77, 83, self._interpretation("bpa2", fallback="atencao"))
        test = service._find_test(self.context, "ravlt")
        if test:
            self._fill_table(5, service._ravlt_rows(test, self.context))
            self._replace(90, 100, self._interpretation("ravlt", fallback="memoria_aprendizagem"))
        test = service._find_test(self.context, "fdt")
        if test:
            self._fill_table(6, service._fdt_rows(test), headers=2, key="wisc_fdt")
            self._replace(103, 112, self._interpretation("fdt", fallback="funcoes_executivas"))
        test = service._find_test(self.context, "etdah_pais")
        if test:
            self._fill_table(7, service._etdah_rows(test), headers=2, key="wisc_etdah_pais")
            self._replace(120, 129, self._interpretation("etdah_pais"))
        self._scared_results()
        test = service._find_test(self.context, "epq_j")
        if test:
            self._fill_table(10, service._epq_rows(test), headers=2, key="wisc_epq")
            self._replace(153, 160, self._interpretation("epq_j"))
        test = service._find_test(self.context, "srs2")
        if test:
            self._fill_table(11, service._srs2_rows(test, "wisc_srs2"), headers=2, key="wisc_srs2")
            self._replace(165, 173, self._interpretation("srs2"))
        # Instrumentos extras continuam disponíveis, mas nunca usam resultados
        # de exemplo do modelo para instrumentos não aplicados.
        if service._find_test(self.context, "etdah_ad"):
            self._extra_etdah_ad()

    def _scared_results(self):
        tests = self.service._find_tests(self.context, "scared")
        if not tests:
            return
        interpretations = []
        for form, table_index in (("child", 8), ("parent", 9)):
            test = next((item for item in tests if ((item.get("classified_payload") or {}).get("form_type") or "child") == form), None)
            self._fill_table(table_index, self.service._scared_rows(test) if test else None, headers=2, key=f"wisc_scared_{'parent' if form == 'parent' else 'self'}")
            if test:
                interpretations.append(self.service._resolve_interpretation_text(None, None, test, self.context))
        self._replace(135, 148, "\n\n".join(interpretations))
        labels = list(dict.fromkeys(self.service._scared_form_label(test) for test in tests))
        description = " e ".join(labels)
        for index, suffix in ((134, f"SCARED – Resultados {description}"), (149, f"SCARED Resultados {description}")):
            found = False
            for run in self.paragraphs[index].runs:
                if "SCARED" in run.text and not found:
                    run.text = re.sub(r"SCARED.*", suffix, run.text)
                    found = True
                elif found:
                    run.text = ""

    def _additional_results(self):
        """Não omitir escalas aplicadas que não têm seção própria no modelo."""
        service = self.service
        anchor = self.paragraphs[175]._p
        body = self.document._body._element

        def text(value, heading=False):
            for line in str(value or "").splitlines():
                if not line.strip():
                    continue
                element = deepcopy(self.paragraphs[150 if heading else 32]._p)
                self._text(Paragraph(element, self.document._body), line.strip())
                anchor.addprevious(element)

        def table(rows, key):
            before = set(body)
            service._append_table_with_interpretation(self.document, rows, key, None)
            for element in list(body):
                if element not in before:
                    anchor.addprevious(element)

        def chart(title, image):
            if not image:
                return
            before = set(body)
            service._append_chart(self.document, title, image)
            for element in list(body):
                if element not in before:
                    anchor.addprevious(element)

        for code, title, loader, key, image in (
            ("bai", "BAI", service._bai_rows, "bai_scores", service._bai_chart),
            ("ebadep_a", "EBADEP-A", service._ebadep_rows, "wisc_ebadep", None),
            ("bfp", "BFP – BATERIA FATORIAL DE PERSONALIDADE", service._bfp_rows, "bfp", service._bfp_chart),
        ):
            test = service._find_test(self.context, code)
            if not test:
                continue
            text(title, heading=True)
            rows = loader(test)
            for values in (rows or []) if code == "bfp" else [rows]:
                table(values, key)
            text(self._interpretation(code, fallback="aspectos_emocionais_comportamentais"))
            if image:
                chart(f"{title} - Resultados", image(test))
        test = service._find_test(self.context, "wasi")
        if test:
            text("WASI – Escala Wechsler Abreviada de Inteligência", heading=True)
            text(service._wasi_intro_text(test, self.context))
            chart("WASI - Resultados", service._wasi_chart_image(test))
            text(self._interpretation("wasi", section="eficiencia_intelectual"))
            for scale in ("verbal", "execucao"):
                table(service._wasi_subscale_rows(test, scale), f"wasi_{scale}")

    def _extra_etdah_ad(self):
        # O modelo não contém E-TDAH-AD: usar uma cópia da tabela/gráfico Pais,
        # com dados e recursos independentes, mantendo sua linguagem visual.
        from docx.opc.part import Part

        service = self.service
        test = service._find_test(self.context, "etdah_ad")
        anchor = next(paragraph._p for paragraph in self.paragraphs[131:176] if paragraph._p.getparent() is self.document._body._element)
        heading = deepcopy(self.paragraphs[116]._p)
        self._text(Paragraph(heading, self.document._body), "E-TDAH-AD")
        anchor.addprevious(heading)
        # Carregar uma cópia limpa, mesmo quando Pais não foi aplicado.
        original = Document(service.WISC4_TEMPLATE_PATH)
        table = deepcopy(original.tables[7]._tbl)
        anchor.addprevious(table)
        from docx.table import Table

        self.tables.append(Table(table, self.document._body))
        self._fill_table(len(self.tables) - 1, service._etdah_rows(test), headers=2, key="wisc_etdah_ad")
        title = self.tables[-1].rows[0].cells[0].paragraphs[0]
        self._text(title, "E-TDAH-AD")
        paragraph = deepcopy(self.paragraphs[120]._p)
        self._text(Paragraph(paragraph, self.document._body), self._interpretation("etdah_ad"))
        anchor.addprevious(paragraph)
        block = deepcopy(original.paragraphs[129]._p)
        for chart in block.iter(qn("c:chart")):
            source = self.document.part.related_parts[chart.get(qn("r:id"))]
            package = self.document.part.package
            root = etree.fromstring(source.blob)
            service._update_chart_title(root, "E-TDAH-AD")
            part = Part(package.next_partname("/word/charts/chart%d.xml"), source.content_type, service._xml_bytes(root), package)
            for rel in source.rels.values():
                if not rel.is_external:
                    part.rels.add_relationship(rel.reltype, rel.target_part, rel.rId)
            chart.set(qn("r:id"), self.document.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"))
        anchor.addprevious(service._bind_chart_block(block, "etdah_ad"))

    def _bind_charts(self):
        service = self.service
        chart_keys = {1: "wisc4", 2: "bpa2", 3: "ravlt", 4: "fdt_auto", 5: "fdt_control", 6: "etdah_pais", 9: "epq_j:raw", 10: "epq_j:percentile", 11: "srs2"}
        scared = service._find_tests(self.context, "scared")
        for form, number in (("child", 7), ("parent", 8)):
            for index, test in enumerate(scared):
                if ((test.get("classified_payload") or {}).get("form_type") or "child") == form:
                    chart_keys[number] = f"scared:{index}"
                    break
        for drawing in list(self.document._body._element.iter(qn("w:drawing"))):
            chart = next(drawing.iter(qn("c:chart")), None)
            if chart is None:
                continue
            properties = next(drawing.iter(qn("wp:docPr")), None)
            if properties is not None and properties.get("descr") == "neuropsi:test:etdah_ad":
                continue
            part = self.document.part.related_parts[chart.get(qn("r:id"))]
            number = int(re.search(r"chart(\d+)\.xml$", str(part.partname)).group(1))
            key = chart_keys.get(number)
            if key is None:
                drawing.getparent().remove(drawing)
            elif properties is not None:
                properties.set("descr", f"neuropsi:test:{key}")

    def _renumber_captions(self):
        counts = {"Tabela": 0, "Gráfico": 0}
        for paragraph in self.document.paragraphs:
            match = re.match(r"\s*(Tabela|Gráfico)\s*\d+", paragraph.text)
            if not match:
                continue
            kind = match.group(1)
            counts[kind] += 1
            # Mantém campos SEQ e seu formato, trocando só o resultado em cache.
            for run in paragraph.runs:
                if run.text.strip().isdigit():
                    run.text = str(counts[kind])
                    break
            else:
                self._text(paragraph, re.sub(r"^(\s*(?:Tabela|Gráfico)\s*)\d+", rf"\g<1>{counts[kind]}", paragraph.text))
