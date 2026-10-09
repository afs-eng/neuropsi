from copy import deepcopy
from datetime import date
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.part import Part
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from apps.reports.builders.references_builder import build_references
from apps.reports.builders.saved_scale_results import numeric_profile, saved_scale_rows
from apps.reports.builders.srs2_profiles import comparable_profiles, profile_label
from apps.reports.builders.wais3_report_builder import WAIS3ReportBuilder
from apps.reports.services.wais3_standardization import WAIS3StandardizationService


class WAIS3DocxBuilder:
    REGULATORY_REFERENCES = (
        "CONSELHO FEDERAL DE PSICOLOGIA. Resolução CFP nº 06/2019. Institui regras para a elaboração de documentos escritos produzidos pela(o) psicóloga(o) no exercício profissional.",
        "CONSELHO FEDERAL DE PSICOLOGIA. Resolução CFP nº 31/2022. Estabelece diretrizes para a realização de Avaliação Psicológica e regulamenta o SATEPSI.",
    )
    EXTRA_REFERENCES = {
        "bfp": "NUNES, C. H. S. S.; HUTZ, C. S.; NUNES, M. F. O. BFP – Bateria Fatorial de Personalidade: manual técnico. São Paulo: Casa do Psicólogo, 2010.",
        "bdefs": "BARKLEY, R. A. BDEFS – Escala de Avaliação de Disfunções Executivas de Barkley. Adaptação brasileira de Victor Polignano Godoy, Leandro Fernandes Malloy-Diniz e Paulo Mattos. São Paulo: Hogrefe, 2018.",
        "iphexa": "LIMA, F. F. V.; SILVA, F. C. IPHEXA – Inventário de Personalidade Hexadimensional. São Paulo: Vetor Editora, 2023.",
    }
    FINAL_CONSIDERATIONS = (
        "A Avaliação Neuropsicológica constitui instrumento para subsidiar o planejamento clínico e orientar estratégias de manejo. "
        "Por se tratar de um documento sigiloso, seu conteúdo só deve ser compartilhado, reproduzido ou discutido mediante autorização "
        "do paciente ou de seu responsável legal. Recomenda-se que o laudo seja lido em sua totalidade, e não apenas em sua conclusão, "
        "para compreensão integral do raciocínio clínico. O uso destas informações deve preservar a integridade socioemocional do paciente, "
        "evitando exposição, discriminação ou constrangimento."
    )
    DOCUMENT_NOTICES = (
        "Não deve ser utilizado para fins diferentes daqueles especificados na identificação do documento.",
        "Possui caráter sigiloso e extrajudicial; sua guarda, compartilhamento e utilização devem observar a finalidade para a qual foi elaborado.",
        "As informações e conclusões devem ser compreendidas de forma integrada aos dados clínicos, à história do desenvolvimento e às observações realizadas.",
        "Esta avaliação observa os princípios éticos profissionais e as Resoluções CFP nº 06/2019 e nº 31/2022, bem como o Código de Ética Profissional do Psicólogo.",
    )

    def __init__(self, service, report, context, sections):
        self.service = service
        self.report = report
        self.context = context
        self.sections = sections
        self.edited_keys = {section.key for section in report.sections.all() if getattr(section, "content_edited", "")}
        self.document = Document(service.WAIS3_STANDARD_TEMPLATE_PATH)
        self.chart_blocks = {
            next(block.iter(qn("wp:docPr"))).get("descr").removeprefix("neuropsi:test:"): block
            for block in service._extract_template_chart_blocks(self.document)
        }
        service._clear_document_body(self.document)
        self.table_index = 0
        self.chart_index = 0

    def heading(self, text, level=1):
        paragraph = self.document.add_paragraph(text, style=f"Heading {level}")
        paragraph.paragraph_format.keep_with_next = True
        return paragraph

    def text(self, text, interpretation=False):
        if interpretation:
            text = self.service._normalize_interpretation_text(text or "")
        for line in (text or "").splitlines():
            if not line.strip():
                continue
            paragraph = self.document.add_paragraph()
            self.service._append_body_text_with_bold_label(paragraph, line.strip())
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.first_line_indent = Cm(1.5)
            paragraph.paragraph_format.space_after = Pt(6)

    def identity(self, label, value):
        self.service._append_label_value(self.document, label, value or "Não informado")

    def table(self, rows, key, caption):
        if not rows:
            return
        self.table_index += 1
        self.service._append_table_with_interpretation(
            self.document, rows, key, None,
            caption=self.service._table_caption_text(self.table_index, caption),
        )

    def chart(self, key, caption, image=None, binding=None):
        block = self.chart_blocks.get(key)
        if block is not None:
            block = self.service._bind_chart_block(block, binding or key)
            if key.startswith("srs2"):
                block = deepcopy(block)
                for chart in block.iter(qn("c:chart")):
                    source = self.document.part.related_parts[chart.get(qn("r:id"))]
                    package = self.document.part.package
                    part = Part(package.next_partname("/word/charts/chart%d.xml"), source.content_type, source.blob, package)
                    for rel in source.rels.values():
                        if not rel.is_external:
                            part.rels.add_relationship(rel.reltype, rel.target_part, rel.rId)
                    chart.set(qn("r:id"), self.document.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart"))
            self.service._append_body_element_before_sectpr(self.document, block)
        elif image:
            self.service._append_chart(self.document, None, image, width=Cm(15.7))
        else:
            return
        self.chart_index += 1
        paragraph = self.document.add_paragraph(self.service._chart_caption_text(self.chart_index, caption))
        self.service._format_caption_paragraph(paragraph)

    def interpretation(self, test, section=None):
        content = self.sections.get(section) if section else None
        return content or self.service._resolve_interpretation_text(None, None, test, self.context)

    def build(self):
        service = self.service
        patient = self.context.get("patient") or {}
        evaluation = self.context.get("evaluation") or {}
        service._add_center_title(self.document, "LAUDO DE AVALIAÇÃO NEUROPSICOLÓGICA")
        self.document.paragraphs[-1].paragraph_format.line_spacing = 1.15
        subtitle = self.document.add_paragraph("Elaborado em conformidade com a Resolução CFP nº 06/2019")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.first_line_indent = Pt(0)
        subtitle.paragraph_format.line_spacing = 1.15
        subtitle.runs[0].font.size = Pt(9)
        self.heading("IDENTIFICAÇÃO")
        self.heading("1.1. Identificação do laudo:", 2)
        self.identity("Autora", service.FIXED_AUTHOR)
        interested = getattr(self.report, "interested_party", "") or ("Familiares" if service._is_adolescent_context(self.context, self.report) else "O paciente" if patient.get("sex") != "F" else "A paciente")
        self.identity("Interessado", interested)
        self.identity("Finalidade", service.FIXED_PURPOSE)
        self.heading("1.2. Identificação do paciente:", 2)
        for label, value in (
            ("Nome", patient.get("full_name")), ("Sexo", service._format_sex_display(patient.get("sex"))),
            ("Data de nascimento", service._format_date_display(patient.get("birth_date"))),
            ("Idade", service._age_text(self.context)), ("Escolaridade", service._schooling_text(patient)),
        ):
            self.identity(label, value)
        if patient.get("mother_name") or patient.get("father_name"):
            self.identity("Filiação", service._filiation_text(patient))
        self.heading("DESCRIÇÃO DA DEMANDA")
        self.heading("Motivo do Encaminhamento", 2)
        self.text(self.sections.get("descricao_demanda") or evaluation.get("referral_reason"))
        self.heading("PROCEDIMENTOS")
        self.text(service._wais3_procedures_intro_text(self.report, self.sections))
        for item in service._adolescent_instruments(self.context):
            service._append_procedure_bullet(self.document, item["name"], item["description"])
        self.heading("ANÁLISE")
        self.text(self.sections.get("historia_pessoal"))
        self.heading("ANÁLISE QUALITATIVA")
        self.intellectual_results()
        self.instrument_results()
        self.heading("CONCLUSÃO")
        self.text(self.sections.get("conclusao") or self.sections.get("hipotese_diagnostica"))
        self.heading("SUGESTÕES DE CONDUTA (ENCAMINHAMENTOS):")
        self.text(self.sections.get("sugestoes_conduta") or self.sections.get("conduta_encaminhamentos"))
        self.heading("CONSIDERAÇÕES FINAIS")
        self.text(self.sections.get("consideracoes_finais") or self.FINAL_CONSIDERATIONS)
        signature_date = evaluation.get("end_date") or date.today().isoformat()
        self.text(f"Goiânia, {service._format_date_display(str(signature_date))}")
        for line in ("Jacqueline O. Caires", "Neuropsicóloga - CRP 09/6017", "Analista do Comportamento", "Especialista em Saúde Mental"):
            paragraph = self.document.add_paragraph(line)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.text("Importante ressaltar que este documento:")
        for index, notice in enumerate(self.DOCUMENT_NOTICES, 1):
            self.text(f"{index}. {notice}")
        self.heading("REFERÊNCIAS BIBLIOGRÁFICAS")
        tests = self.context.get("validated_tests") or []
        references = [*self.REGULATORY_REFERENCES, *build_references(tests)[1:]]
        codes = {test.get("instrument_code") for test in tests}
        references.extend(reference for code, reference in self.EXTRA_REFERENCES.items() if code in codes)
        self.text(self.sections.get("referencias_bibliograficas") or "\n".join(dict.fromkeys(references)))
        return self.document

    def intellectual_results(self):
        service = self.service
        test = service._find_test(self.context, "wais3")
        self.text(service._wais3_intro_text(test, self.context))
        descriptions = (
            "Avalia o conhecimento adquirido, o raciocínio verbal e a formação de conceitos.",
            "Avalia raciocínio não verbal, organização perceptual e análise visuoespacial.",
            "Avalia atenção, concentração e manipulação mental de informações.",
            "Avalia rapidez, precisão e eficiência visuomotora em tarefas automatizadas.",
            "Avalia compreensão, expressão verbal e raciocínio mediado pela linguagem.",
            "Avalia raciocínio não verbal e resolução de problemas visuoespaciais.",
            "Estima habilidades intelectuais gerais com menor influência da memória operacional e da velocidade de processamento.",
        )
        for (lead, _), description in zip(service._wais3_global_bullet_parts(test), descriptions):
            self.text(f"{lead}. {description}")
        self.heading("Desempenho do paciente no WAIS III", 2)
        self.chart("wais3", "WAIS III - ÍNDICES DE QIS")
        self.text(self.sections.get("capacidade_cognitiva_global") or service._wais3_compact_interpretation_text(test, self.context), True)
        self.heading("Subescalas WAIS III", 2)
        for block in WAIS3ReportBuilder.for_adult(set()).subtest_blocks:
            self.heading(block.title, 3)
            self.table(service._wais3_domain_rows(test, block.interpretation_section), block.table_key, block.table_caption)
            self.text(self.sections.get(block.interpretation_section) or WAIS3StandardizationService.build(block.interpretation_section, self.context), True)

    def instrument_results(self):
        service = self.service
        for code in ("bpa2", "ravlt", "fdt", "etdah_ad", "bdefs", "ebadep_a", "ebadep_ij", "ebaped_ij", "bfp", "iphexa", "bai", "etdah_pais", "scared", "epq_j"):
            for test in service._find_tests(self.context, code):
                if code == "bpa2":
                    self.heading("BPA-2 Bateria Psicológica para Avaliação da Atenção", 2)
                    self.text("A BPA-2 investiga atenção concentrada, dividida, alternada e atenção geral.")
                    self.table(service._bpa_rows(test, self.context), "bpa", "Atenção BPA-2 Resultados")
                    if self.sections.get("bpa2"):
                        self.text(self.sections["bpa2"], True)
                    else:
                        service._append_bpa_interpretation_block(self.document, test, self.context)
                    self.chart(code, "BPA-2 apresenta os resultados da avaliação da atenção")
                elif code == "ravlt":
                    self.heading("RAVLT Rey Auditory Verbal Learning Test", 2)
                    service._append_ravlt_conceptual_paragraph(self.document)
                    self.chart(code, "RAVLT Resultados")
                    self.table(service._ravlt_rows(test, self.context), "ravlt", service._ravlt_table_caption())
                    self.text(self.interpretation(test, "ravlt"), True)
                elif code == "fdt":
                    self.heading("FDT- TESTE DOS CINCO DÍGITOS", 2)
                    self.text(service._fdt_description_text())
                    self.table(service._fdt_rows(test), "fdt", "FDT Processos Automáticos e Controlados")
                    self.text(self.interpretation(test, "fdt"), True)
                    self.chart("fdt_auto", "FDT Processos Automáticos")
                    self.chart("fdt_control", "FDT Processos Controlados")
                elif code in {"etdah_ad", "etdah_pais"}:
                    self.heading("ETDAH-AD" if code == "etdah_ad" else "E-TDAH-PAIS", 2)
                    self.table(service._etdah_rows(test), "wais_etdah_ad" if code == "etdah_ad" else "etdah_pais", "E-TDAH Resultados")
                    self.text(self.interpretation(test, code), True)
                    self.chart(code, "E-TDAH Resultados", service._etdah_chart(test))
                elif code == "bfp":
                    self.heading("BFP – Bateria Fatorial de Personalidade", 2)
                    self.text(service._bfp_description_text())
                    self.chart(code, "BFP - Radar das facetas – Resultados", service._bfp_chart(test))
                    self.text(self.interpretation(test, code), True)
                    tables = service._bfp_rows(test) or []
                    if tables:
                        self.table([row for rows in tables for row in rows], "bfp", "Resultados gerais")
                elif code in {"ebadep_a", "ebadep_ij", "ebaped_ij", "bai", "scared", "epq_j"}:
                    names = {"bai": "BAI", "scared": f"SCARED — {service._scared_form_label(test)}", "epq_j": "EPQ-J"}
                    self.heading(names.get(code, "EBADEP-A" if code == "ebadep_a" else "EBADEP-IJ"), 2)
                    loader = service._ebadep_rows if code.startswith(("ebadep", "ebaped")) else getattr(service, f"_{'epq' if code == 'epq_j' else code}_rows")
                    table_key = "wais_ebadep" if code.startswith(("ebadep", "ebaped")) else "bai_scores" if code == "bai" else service._scared_table_key(test, self.context) if code == "scared" else "epq"
                    self.table(loader(test), table_key, "Resultados do instrumento")
                    self.text(self.interpretation(test, code), True)
                    if code in {"bai", "scared", "epq_j"}:
                        chart_loader = getattr(service, f"_{'epq' if code == 'epq_j' else code}_chart")
                        self.chart(code, "Resultados do instrumento", chart_loader(test))
                else:
                    self.generic_results(test, code)
        self.srs2_results()

    def generic_results(self, test, code):
        titles = {"bdefs": "BDEFS Escala de Déficits em Funções Executivas", "iphexa": "IPHEXA Inventário de Personalidade Hexadimensional"}
        self.heading(titles.get(code) or test.get("instrument_name") or code.upper(), 2)
        results = saved_scale_rows(test)
        labels, values = numeric_profile(results, "percentile" if code == "iphexa" else "score")
        if code == "iphexa" and labels:
            from apps.reports.charts.saved_personality_chart import saved_percentile_radar

            self.chart(code, "IPHEXA - Radar dos resultados", saved_percentile_radar(labels, values))
        if code == "bdefs" and labels:
            total = [(label, score) for label, score in zip(labels, values) if "total" in label.casefold()]
            domains = [(label, score) for label, score in zip(labels, values) if "total" not in label.casefold()]
            for caption, profile in (("BDEFS - Resultados dos domínios", domains), ("BDEFS - Resultado total", total)):
                if profile:
                    image = self.service._build_chart_png("bar", caption, [label for label, _ in profile], [score for _, score in profile], "Escore bruto")
                    self.chart(code, caption, image)
        self.text(self.interpretation(test, code), True)
        if results:
            rows = [["Domínio", "Escore", "Percentil", "Classificação"]]
            for row in results:
                rows.append([row["label"], self.service._num(row["score"]), self.service._num(row["percentile"]), row["classification"] or "-"])
            self.table(rows, "wais_saved_results", f"{code.upper()} Resultados")
        else:
            self.text("Resultados estruturados não disponíveis para apresentação tabular.")

    def srs2_results(self):
        tests = self.service._find_tests(self.context, "srs2")
        if not tests:
            return
        self.heading("SRS-2 Escala de Responsividade Social", 2)
        self.text(self.service._srs2_description_text())
        if len(tests) > 1 and "srs2" in self.edited_keys:
            self.text(self.sections["srs2"], True)
        for index, test in enumerate(tests):
            label = profile_label(test, index)
            self.heading(label, 3)
            self.table(self.service._srs2_rows(test, "wais_srs2"), "wais_srs2", f"SRS-2 Resultados — {label}")
            self.text(self.interpretation(test, "srs2" if len(tests) == 1 else None), True)
            self.chart("srs2", f"SRS-2 Resultados — {label}", binding="srs2" if len(tests) == 1 else f"srs2:{index}")
        categories, scores, labels = comparable_profiles(tests)
        if categories:
            self.heading("Análise Integrada", 3)
            rows = [["Domínio", *labels]] + [[name, *(self.service._num(series[index]) for series in scores)] for index, name in enumerate(categories)]
            self.table(rows, "wais_comparison", "SRS-2 Comparação dos escores T entre respondentes")
            self.chart("srs2_comparison", "SRS-2 Resultados das discrepâncias")
            self.text("Em análise clínica, as convergências e discrepâncias entre respondentes devem ser integradas à história do desenvolvimento, à observação clínica e ao funcionamento cotidiano. A comparação descreve diferenças de percepção; a SRS-2, isoladamente, não estabelece diagnóstico.")
