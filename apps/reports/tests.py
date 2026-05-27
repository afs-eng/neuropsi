from django.test import SimpleTestCase
from xml.etree import ElementTree as ET
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from datetime import date
from unittest.mock import patch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from lxml import etree as LET

from apps.reports.builders.tests_builder import _build_wais3_tables
from apps.reports.builders.references_builder import build_references
from apps.reports.services.report_export_service import ReportExportService
from apps.reports.services.report_generation_service import ReportGenerationService
from apps.reports.services.section_regeneration_service import SectionRegenerationService
from apps.reports.services.wisc4_standardization import WISC4StandardizationService


class ReferencesBuilderTests(SimpleTestCase):
    def test_build_references_includes_only_validated_tests(self):
        references = build_references(
            [
                {"instrument_code": "wisc4"},
                {"instrument_code": "fdt"},
            ]
        )

        self.assertEqual(len(references), 3)
        self.assertIn("WISC-IV", references[1])
        self.assertIn("Teste dos Cinco Dígitos", references[2])

    def test_build_references_deduplicates_equivalent_entries(self):
        references = build_references(
            [
                {"instrument_code": "ebadep_ij"},
                {"instrument_code": "ebaped_ij"},
            ]
        )

        self.assertEqual(len(references), 2)
        self.assertIn("EBADEP-IJ", references[1])


class WISC4StandardizationTests(SimpleTestCase):
    def setUp(self):
        self.context = {
            "patient": {"full_name": "Maria Clara", "sex": "F"},
            "validated_tests": [
                {
                    "instrument_code": "wisc4",
                    "structured_results": {
                        "qit_data": {
                            "escore_composto": 61,
                            "classificacao": "Extremamente Baixo",
                        },
                        "gai_data": {
                            "escore_composto": 70,
                            "classificacao": "Limítrofe",
                        },
                        "cpi_data": {
                            "escore_composto": 58,
                            "classificacao": "Extremamente Baixo",
                        },
                        "indices": [
                            {
                                "indice": "icv",
                                "escore_composto": 84,
                                "classificacao": "Média Inferior",
                            },
                            {
                                "indice": "iop",
                                "escore_composto": 61,
                                "classificacao": "Extremamente Baixo",
                            },
                            {
                                "indice": "imt",
                                "escore_composto": 52,
                                "classificacao": "Extremamente Baixo",
                            },
                            {
                                "indice": "ivp",
                                "escore_composto": 68,
                                "classificacao": "Extremamente Baixo",
                            },
                        ],
                        "subtestes": [
                            {"codigo": "SM", "classificacao": "Média"},
                            {"codigo": "CN", "classificacao": "Média Inferior"},
                            {"codigo": "CO", "classificacao": "Média"},
                            {"codigo": "RM", "classificacao": "Limítrofe"},
                            {"codigo": "VC", "classificacao": "Limítrofe"},
                            {"codigo": "CB", "classificacao": "Extremamente Baixo"},
                            {"codigo": "SNL", "classificacao": "Extremamente Baixo"},
                            {"codigo": "DG", "classificacao": "Limítrofe"},
                        ],
                    },
                }
            ],
        }

    def test_build_global_section_uses_standardized_wisc_template(self):
        text = WISC4StandardizationService.build(
            "capacidade_cognitiva_global", self.context
        )

        self.assertIn("Capacidade Cognitiva Global: a avaliação neuropsicológica de Maria", text)
        self.assertIn("Quociente de Inteligência Total (QIT = 61)", text)
        self.assertIn("Compreensão Verbal (ICV) — 84 — Média Inferior", text)
        self.assertNotIn("Índice de Habilidade Geral (GAI = 70)", text)
        self.assertNotIn("O Índice de Compreensão Verbal", text)

    def test_build_global_section_includes_fixed_cognitive_age_phrase_when_available(self):
        self.context["validated_tests"][0]["structured_results"]["idade_cognitiva"] = "12 anos e 6 meses"

        text = WISC4StandardizationService.build(
            "capacidade_cognitiva_global", self.context
        )

        self.assertIn(
            "com idade cognitiva estimada de 12 anos e 6 meses",
            text,
        )

    def test_build_domain_section_uses_standardized_wisc_subscale_template(self):
        text = WISC4StandardizationService.build("funcoes_executivas", self.context)

        self.assertIn("A avaliação das funções executivas de Maria", text)
        self.assertIn("No subteste Semelhanças", text)
        self.assertIn("No subteste Raciocínio Matricial", text)
        self.assertIn("perfil executivo heterogêneo", text)


class SectionRegenerationServiceTests(SimpleTestCase):
    def test_enabled_test_section_keys_follow_validated_tests(self):
        context = {
            "validated_tests": [
                {"instrument_code": "bpa2"},
                {"instrument_code": "fdt"},
                {"instrument_code": "ebadep_ij"},
            ]
        }

        section_keys = SectionRegenerationService._enabled_test_section_keys(context)

        self.assertIn("bpa2", section_keys)
        self.assertIn("fdt", section_keys)
        self.assertIn("ebadep", section_keys)
        self.assertNotIn("ravlt", section_keys)

    @patch("apps.reports.services.section_regeneration_service.build_report_snapshot")
    def test_regenerate_test_sections_creates_missing_sections_with_scope(
        self, build_snapshot
    ):
        build_snapshot.return_value = {"validated_tests": [{"instrument_code": "bpa2"}]}

        class SectionQuerySet(list):
            def all(self):
                return self

            def filter(self, **kwargs):
                key = kwargs.get("key")
                return SectionQuerySet([item for item in self if item.key == key])

            def first(self):
                return self[0] if self else None

        class Section:
            def __init__(self, key, title, order, is_locked=False):
                self.key = key
                self.title = title
                self.order = order
                self.is_locked = is_locked
                self.generation_metadata = {}

            def save(self, update_fields=None):
                return None

            def delete(self):
                sections.remove(self)

        class Sections(SectionQuerySet):
            def create(self, **kwargs):
                section = Section(**kwargs)
                self.append(section)
                return section

        sections = Sections()

        class Report:
            def __init__(self):
                self.evaluation = object()
                self.sections = sections
                self.status = "draft"
                self.context_payload = {}
                self.ai_metadata = {}
                self.generated_at = None

            def save(self, update_fields=None):
                return None

        report = Report()

        with patch.object(
            SectionRegenerationService,
            "_regenerate_section_in_place",
            side_effect=lambda report, section_key, context, ensure_ai_available=False: report.sections.filter(key=section_key).first(),
        ), patch(
            "apps.reports.services.section_regeneration_service.ReportSectionService._rebuild_report_text"
        ), patch(
            "apps.reports.services.section_regeneration_service.ReportVersionService.create_version"
        ):
            regenerated = SectionRegenerationService.regenerate_test_sections(
                report,
                scope="tests_only",
            )

        self.assertEqual(regenerated, ["bpa2"])
        self.assertEqual(report.ai_metadata["generation_scope"], "tests_only")
        self.assertEqual([section.key for section in report.sections.all()], ["bpa2"])


class ReportGenerationServiceConclusionTests(SimpleTestCase):
    def test_wisc4_adolescent_conclusion_follows_fixed_model_order(self):
        class Section:
            def __init__(self, key, generated):
                self.key = key
                self.content_edited = ""
                self.content_generated = generated

        class Sections(list):
            def all(self):
                return self

            def order_by(self, _field):
                return self

            def filter(self, **kwargs):
                key = kwargs.get("key")
                return Sections([item for item in self if item.key == key])

            def first(self):
                return self[0] if self else None

        class Report:
            sections = Sections(
                [
                    Section("capacidade_cognitiva_global", "Perfil heterogêneo com fragilidades cognitivas."),
                    Section("funcoes_executivas", "Fragilidades executivas e atencionais com oscilação clínica."),
                    Section("memoria_aprendizagem", "Perfil heterogêneo na memória e aprendizagem."),
                    Section("linguagem", "Recursos de linguagem relativamente preservados."),
                    Section("gnosias_praxias", "Fragilidades visuoconstrutivas pontuais."),
                    Section("aspectos_emocionais_comportamentais", "Indicadores emocionais e comportamentais relevantes."),
                    Section("srs2", "Indicadores sociais que exigem correlação clínica."),
                ]
            )

        context = {
            "patient": {"full_name": "Maria Clara", "sex": "F", "schooling": "fundamental", "birth_date": "2010-01-10"},
            "evaluation": {"clinical_hypothesis": "Transtorno do Déficit de Atenção e Hiperatividade (TDAH), apresentação combinada"},
            "validated_tests": [
                {
                    "instrument_code": "wisc4",
                    "structured_results": {
                        "qit_data": {"escore_composto": 61, "classificacao": "Extremamente Baixo"},
                        "indices": [
                            {"indice": "icv", "escore_composto": 84, "classificacao": "Média Inferior"},
                            {"indice": "iop", "escore_composto": 61, "classificacao": "Extremamente Baixo"},
                            {"indice": "imt", "escore_composto": 52, "classificacao": "Extremamente Baixo"},
                            {"indice": "ivp", "escore_composto": 68, "classificacao": "Extremamente Baixo"},
                        ],
                    },
                },
                {"instrument_code": "srs2"},
            ],
        }
        context["evaluation"]["start_date"] = "2024-04-20"

        text = ReportGenerationService._build_conclusion_text(Report(), context)

        self.assertIn("Maria Clara apresenta funcionamento neuropsicológico", text)
        self.assertIn("Na Escala Wechsler de Inteligência para Crianças – Quarta Edição (WISC-IV)", text)
        self.assertIn("No domínio atencional e executivo", text)
        self.assertIn("A avaliação da memória e aprendizagem", text)
        self.assertIn("Nos domínios de linguagem, gnosias e praxias", text)
        self.assertIn("Nos aspectos emocionais e comportamentais", text)
        self.assertIn("Na avaliação do funcionamento social", text)
        self.assertIn("Diante da integração dos resultados das testagens", text)
        self.assertIn("Hipótese Diagnóstica:", text)
        self.assertIn("Ressalta-se que o ser humano possui natureza dinâmica", text)
        self.assertNotIn("verifica-se", text.casefold())


class WISC4ExportTableTests(SimpleTestCase):
    def setUp(self):
        self.context = {
            "patient": {"birth_date": "2016-01-10"},
            "evaluation": {"start_date": "2024-04-20"},
            "validated_tests": [
                {
                    "instrument_code": "wisc4",
                    "applied_on": "2024-04-20",
                }
            ],
        }
        self.test = {
            "classified_payload": {
                "subtestes": [
                    {
                        "codigo": "SM",
                        "subteste": "Semelhanças",
                        "escore_bruto": 23,
                        "classificacao": "Média",
                    },
                    {
                        "codigo": "CN",
                        "subteste": "Conceitos Figurativos",
                        "escore_bruto": 11,
                        "classificacao": "Média Inferior",
                    },
                    {
                        "codigo": "CO",
                        "subteste": "Compreensão",
                        "escore_bruto": 24,
                        "classificacao": "Média",
                    },
                    {
                        "codigo": "RM",
                        "subteste": "Raciocínio Matricial",
                        "escore_bruto": 12,
                        "classificacao": "Limítrofe",
                    },
                ]
            }
        }

    def test_wisc_reference_scores_use_age_norm_table_ranges(self):
        table = ReportExportService._wisc_ncp_table(self.context)

        self.assertEqual(
            ReportExportService._wisc_reference_scores(table, "SM"),
            ("32-44", "6-16", "1"),
        )

    def test_wisc_rows_build_expected_columns_for_executive_domain(self):
        rows = ReportExportService._wisc_rows(
            self.test, self.context, "funcoes_executivas"
        )

        self.assertEqual(
            rows[0],
            [
                "Testes Utilizados",
                "Escore Máximo",
                "Escore Médio",
                "Escore Mínimo",
                "Escore Bruto",
                "Classificação",
            ],
        )
        self.assertEqual(rows[1], ["Semelhanças", "32-44", "6-16", "1", "23", "Média"])


class WAIS3ExportTableTests(SimpleTestCase):
    def _evaluation(self):
        class Patient:
            birth_date = date(2000, 1, 1)

        class Evaluation:
            patient = Patient()
            start_date = date(2020, 6, 1)

        return Evaluation()

    def test_build_wais3_tables_uses_normative_ranges_and_spontaneous_speech(self):
        payload = {
            "subtestes": {
                "semelhancas": {"nome": "Semelhanças", "pontos_brutos": 18, "classificacao": "Média"},
                "vocabulario": {"nome": "Vocabulário", "pontos_brutos": 28, "classificacao": "Média"},
                "compreensao": {"nome": "Compreensão", "pontos_brutos": 16, "classificacao": "Média"},
            }
        }

        tables = _build_wais3_tables(payload, self._evaluation(), date(2020, 6, 1))

        self.assertEqual(tables["linguagem"][0], {
            "label": "Semelhanças",
            "maxScore": "38",
            "avgScore": "17-28",
            "minScore": "9-10",
            "obtainedScore": "18",
            "classification": "Média",
        })
        self.assertEqual(tables["linguagem"][1], {
            "label": "Vocabulário",
            "maxScore": "66",
            "avgScore": "23-42",
            "minScore": "11-14",
            "obtainedScore": "28",
            "classification": "Média",
        })
        self.assertEqual(tables["linguagem"][2], {
            "label": "Compreensão",
            "maxScore": "33",
            "avgScore": "13-26",
            "minScore": "5-6",
            "obtainedScore": "16",
            "classification": "Média",
        })
        self.assertEqual(tables["linguagem"][3], {
            "label": "Fala Espontânea",
            "note": "Fala espontânea dentro do esperado para a sua idade",
        })

    def test_build_wais3_tables_recomputes_classification_from_normative_table(self):
        payload = {
            "subtestes": {
                "vocabulario": {
                    "nome": "Vocabulário",
                    "pontos_brutos": 39,
                    "classificacao": "Média Superior",
                },
            }
        }

        tables = _build_wais3_tables(payload, self._evaluation(), date(2020, 6, 1))

        self.assertEqual(tables["linguagem"][0], {
            "label": "Vocabulário",
            "maxScore": "66",
            "avgScore": "23-42",
            "minScore": "11-14",
            "obtainedScore": "39",
            "classification": "Média",
        })

    def test_build_wais3_tables_uses_normative_classification_across_domains(self):
        payload = {
            "subtestes": {
                "semelhancas": {"nome": "Semelhanças", "pontos_brutos": 28, "classificacao": "Média Superior"},
                "vocabulario": {"nome": "Vocabulário", "pontos_brutos": 42, "classificacao": "Média Superior"},
                "compreensao": {"nome": "Compreensão", "pontos_brutos": 26, "classificacao": "Média Superior"},
                "raciocinio_matricial": {"nome": "Raciocínio Matricial", "pontos_brutos": 21, "classificacao": "Média Superior"},
                "cubos": {"nome": "Cubos", "pontos_brutos": 44, "classificacao": "Média Superior"},
                "sequencia_numeros_letras": {"nome": "Sequência de Números e Letras", "pontos_brutos": 11, "classificacao": "Média Superior"},
                "digitos": {"nome": "Dígitos", "pontos_brutos": 17, "classificacao": "Média Superior"},
            }
        }

        tables = _build_wais3_tables(payload, self._evaluation(), date(2020, 6, 1))

        expected = {
            "linguagem": {
                "Semelhanças": "Média",
                "Vocabulário": "Média",
                "Compreensão": "Média",
            },
            "gnosias_praxias": {
                "Raciocínio Matricial": "Média",
                "Cubos": "Média",
            },
            "funcoes_executivas": {
                "Semelhanças": "Média",
                "Compreensão": "Média",
                "Raciocínio Matricial": "Média",
            },
            "memoria_aprendizagem": {
                "Sequência de Números e Letras": "Média",
                "Dígitos": "Média",
                "RAVLT": "Leitura do Gráfico",
            },
        }

        for domain, rows in tables.items():
            for row in rows:
                if row.get("note"):
                    continue
                self.assertEqual(row["classification"], expected[domain][row["label"]])

    def test_build_wais3_tables_keeps_non_average_normative_classification(self):
        payload = {
            "subtestes": {
                "vocabulario": {"nome": "Vocabulário", "pontos_brutos": 11, "classificacao": "Média Superior"},
                "cubos": {"nome": "Cubos", "pontos_brutos": 45, "classificacao": "Média"},
            }
        }

        tables = _build_wais3_tables(payload, self._evaluation(), date(2020, 6, 1))

        self.assertEqual(tables["linguagem"][0]["classification"], "Limítrofe")
        self.assertEqual(tables["gnosias_praxias"][0]["classification"], "Média Superior")

    def test_wais3_domain_rows_use_dynamic_template_structure(self):
        test = {
            "wais3_tables": {
                "linguagem": [
                    {
                        "label": "Semelhanças",
                        "maxScore": "38",
                        "avgScore": "17-28",
                        "minScore": "9-10",
                        "obtainedScore": "18",
                        "classification": "Média",
                    },
                    {
                        "label": "Fala Espontânea",
                        "note": "Fala espontânea dentro do esperado para a sua idade",
                    },
                ]
            }
        }

        rows = ReportExportService._wais3_domain_rows(test, "linguagem")

        self.assertEqual(rows[0], ["Testes Utilizados", "Escore Máximo", "Escore Médio", "Escore Mínimo", "Escore Bruto", "Classificação"])
        self.assertEqual(rows[1], ["Semelhanças", "38", "17-28", "9-10", "18", "Média"])
        self.assertEqual(rows[2], ["Fala Espontânea", "Fala espontânea dentro do esperado para a sua idade", "", "", "", ""])

    def test_build_wais3_tables_adds_ravlt_placeholder_to_memory_domain(self):
        payload = {
            "subtestes": {
                "sequencia_numeros_letras": {"nome": "Sequência de Números e Letras", "pontos_brutos": 11, "classificacao": "Média"},
                "digitos": {"nome": "Dígitos", "pontos_brutos": 17, "classificacao": "Média"},
            }
        }

        tables = _build_wais3_tables(payload, self._evaluation(), date(2020, 6, 1))

        self.assertEqual(tables["memoria_aprendizagem"][-1], {
            "label": "RAVLT",
            "maxScore": "-",
            "avgScore": "-",
            "minScore": "-",
            "obtainedScore": "-",
            "classification": "Leitura do Gráfico",
        })


class ReportExportChartSanitizationTests(SimpleTestCase):
    def test_etdah_table_title_distinguishes_ad_from_pais(self):
        self.assertEqual(ReportExportService._table_title_text("etdah_ad"), "E-TDAH-AD")
        self.assertEqual(ReportExportService._table_title_text("etdah_pais"), "E-TDAH-PAIS")

    def test_wisc4_model_titles_and_captions_are_canonical(self):
        self.assertEqual(
            ReportExportService._table_title_text("scared_parent"),
            "SCARED - Mãe",
        )
        self.assertEqual(
            ReportExportService._table_title_text("epq_j"),
            "EPQ-J - Inventário de Personalidade de Eysenck para Jovens",
        )
        self.assertEqual(
            ReportExportService._canonical_table_caption("etdah_ad", "qualquer"),
            "E-TDAH-AD Resultados",
        )

    def test_wisc4_model_table_widths_follow_template_measurements(self):
        self.assertEqual(
            [round(width.cm, 2) for width in ReportExportService._table_widths("wisc_linguagem")],
            [3.06, 2.88, 2.58, 2.81, 2.64, 2.49],
        )
        self.assertEqual(
            [round(width.cm, 2) for width in ReportExportService._table_widths("etdah_pais")],
            [7.16, 2.8, 1.51, 1.92, 3.1],
        )

    def test_wais3_model_titles_and_widths_are_canonical(self):
        self.assertEqual(
            ReportExportService._table_title_text("wais_etdah_ad"),
            "E-TDAH-AD",
        )
        self.assertEqual(
            ReportExportService._table_title_text("wais_srs2_adulto"),
            "SRS-2 ADULTO AUTORRELATO",
        )
        self.assertEqual(
            [round(width.cm, 2) for width in ReportExportService._table_widths("wais_linguagem")],
            [3.2, 2.7, 2.6, 2.6, 2.5, 2.88],
        )
        self.assertEqual(
            [round(width.cm, 2) for width in ReportExportService._table_widths("wais_etdah_ad")],
            [8.45, 2.45, 1.3, 1.7, 2.58],
        )

    def test_wasi_model_table_widths_and_template_are_canonical(self):
        self.assertEqual(
            ReportExportService.WASI_TEMPLATE_PATH,
            ReportExportService.TEMPLATE_DIR / "MODELO-WASI.docx",
        )
        self.assertEqual(
            [round(width.cm, 2) for width in ReportExportService._table_widths("wasi_verbal")],
            [4.45, 2.41, 2.41, 2.41, 2.29, 2.92],
        )
        self.assertEqual(
            [round(width.cm, 2) for width in ReportExportService._table_widths("wasi_etdah_ad")],
            [9.02, 2.24, 1.23, 1.56, 2.44],
        )

    def test_etdah_ad_rows_follow_official_factor_order(self):
        rows = ReportExportService._etdah_rows(
            {
                "classified_payload": {
                    "results": {
                        "D": {"name": "Fator 1 - Desatenção (D)", "raw_score": 64, "mean": 42.3, "percentile_text": "86", "classification": "Superior"},
                        "H": {"name": "Fator 5 - Hiperatividade (H)", "raw_score": 20, "mean": 14.35, "percentile_text": "85", "classification": "Superior"},
                        "I": {"name": "Fator 2 - Impulsividade (I)", "raw_score": 57, "mean": 44.3, "percentile_text": "78,8", "classification": "Média Superior"},
                        "AE": {"name": "Fator 3 - Aspectos Emocionais (AE)", "raw_score": 6, "mean": 5.77, "percentile_text": "60", "classification": "Média"},
                        "AAMA": {"name": "Fator 4 - Autorregulação da Atenção, Motivação e Ação (AAMA)", "raw_score": 25, "mean": 21.1, "percentile_text": "75", "classification": "Média Superior"},
                    }
                }
            }
        )

        self.assertEqual(
            [row[0] for row in rows[1:]],
            [
                "Fator 1 - Desatenção (D)",
                "Fator 2 - Impulsividade (I)",
                "Fator 3 - Aspectos Emocionais (AE)",
                "Fator 4 - Autorregulação da Atenção, Motivação e Ação (AAMA)",
                "Fator 5 - Hiperatividade (H)",
            ],
        )

    def test_insert_etdah_ad_interpretation_formats_factor_heading_and_body(self):
        document = Document()
        start = document.add_paragraph("ANÁLISE QUALITATIVA")
        interpretation = (
            "Interpretação e Observações Clínicas: **Texto introdutório.**\n\n"
            "No Fator 1 — Desatenção, o resultado classificado como **Média Inferior** indica funcionamento dentro dos limites esperados.\n\n"
            "Dessa forma, **não** há hipótese diagnóstica isolada."
        )

        ReportExportService._insert_etdah_ad_interpretation_block_after(
            start,
            ReportExportService._normalize_interpretation_text(interpretation),
        )

        paragraphs = [p for p in document.paragraphs if p.text.strip()]

        self.assertEqual(paragraphs[1].text.strip(), "Interpretação e Observações Clínicas: Texto introdutório.")
        self.assertEqual(paragraphs[2].text.strip(), "Fator 1 – Desatenção")
        self.assertTrue(any(run.bold for run in paragraphs[2].runs if run.text.strip()))
        self.assertEqual(paragraphs[3].text.strip(), "o resultado classificado como Média Inferior indica funcionamento dentro dos limites esperados.")
        self.assertEqual(paragraphs[3].alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        self.assertAlmostEqual(paragraphs[3].paragraph_format.first_line_indent.cm, 1.5, places=2)
        self.assertEqual(paragraphs[4].text.strip(), "Dessa forma, não há hipótese diagnóstica isolada.")

    def test_wais3_template_is_organized_under_templates_assets(self):
        self.assertEqual(
            ReportExportService.WAIS3_TEMPLATE_PATH,
            ReportExportService.TEMPLATE_DIR / "Modelo-WAIS3.docx",
        )
        self.assertTrue(ReportExportService.WAIS3_TEMPLATE_PATH.exists())

    def test_wisc4_template_is_organized_under_templates_assets(self):
        self.assertEqual(
            ReportExportService.WISC4_TEMPLATE_PATH,
            ReportExportService.TEMPLATE_DIR / "Modelo-WISC4.docx",
        )
        self.assertTrue(ReportExportService.WISC4_TEMPLATE_PATH.exists())

    def test_adult_reports_restore_header_footer_from_papel_timbrado(self):
        class Patient:
            age = 23

        class Report:
            patient = Patient()

        template_path = ReportExportService.WAIS3_TEMPLATE_PATH

        self.assertEqual(
            ReportExportService._header_footer_template_path(
                template_path,
                Report(),
                {"patient": {"birth_date": "2002-09-12"}},
            ),
            ReportExportService.DEFAULT_TEMPLATE_PATH,
        )

    def test_wais3_reports_restore_header_footer_from_own_template(self):
        class Patient:
            age = 23

        class Report:
            patient = Patient()

        template_path = ReportExportService.WAIS3_TEMPLATE_PATH

        self.assertEqual(
            ReportExportService._header_footer_template_path(
                template_path,
                Report(),
                {"patient": {"birth_date": "2002-09-12"}, "validated_tests": [{"instrument_code": "wais3"}]},
            ),
            template_path,
        )

    def _report_stub(self):
        class Report:
            evaluation = None
            interested_party = ""
            purpose = ""

        return Report()

    def test_apply_base_styles_preserves_template_header_footer_layout(self):
        document = Document(str(ReportExportService.WAIS3_TEMPLATE_PATH))
        section = document.sections[0]

        original = {
            "top": section.top_margin,
            "bottom": section.bottom_margin,
            "left": section.left_margin,
            "right": section.right_margin,
            "header": section.header_distance,
            "footer": section.footer_distance,
        }

        ReportExportService._apply_base_styles(document)

        self.assertEqual(section.top_margin, original["top"])
        self.assertEqual(section.bottom_margin, original["bottom"])
        self.assertEqual(section.left_margin, original["left"])
        self.assertEqual(section.right_margin, original["right"])
        self.assertEqual(section.header_distance, original["header"])
        self.assertEqual(section.footer_distance, original["footer"])

    def test_restore_template_header_footer_replaces_exported_header_with_template(self):
        template_bytes = ReportExportService.WAIS3_TEMPLATE_PATH.read_bytes()
        mutated = template_bytes.replace(b"word/media/image4.jpeg", b"word/media/image1.jpeg")

        restored = ReportExportService._restore_template_header_footer(
            mutated,
            ReportExportService.WAIS3_TEMPLATE_PATH,
        )

        with ZipFile(BytesIO(restored), "r") as archive:
            header_xml = archive.read("word/header1.xml")
            header_rels = archive.read("word/_rels/header1.xml.rels")

        with ZipFile(BytesIO(template_bytes), "r") as template_archive:
            expected_header_xml = template_archive.read("word/header1.xml")
            expected_header_rels = template_archive.read("word/_rels/header1.xml.rels")

        self.assertEqual(header_xml, expected_header_xml)
        self.assertEqual(header_rels, expected_header_rels)

    def test_restore_template_header_footer_restores_document_section_references(self):
        template_bytes = ReportExportService.WAIS3_TEMPLATE_PATH.read_bytes()

        with ZipFile(BytesIO(template_bytes), "r") as archive:
            document_xml = archive.read("word/document.xml").decode("utf-8")

        mutated_document_xml = document_xml.replace(
            '<w:sectPr w:rsidR="00F43DEA" w:rsidRPr="00A64CAA" w:rsidSect="00C77ADB"><w:headerReference w:type="default" r:id="rId18"/><w:footerReference w:type="default" r:id="rId19"/><w:pgSz w:w="11906" w:h="16838"/><w:pgMar w:top="1701" w:right="1134" w:bottom="1304" w:left="1418" w:header="397" w:footer="283" w:gutter="0"/><w:cols w:space="720"/><w:formProt w:val="0"/><w:docGrid w:linePitch="326" w:charSpace="8192"/></w:sectPr>',
            '',
        )

        buffer = BytesIO()
        with ZipFile(BytesIO(template_bytes), "r") as source_zip, ZipFile(buffer, "w") as target_zip:
            for item in source_zip.infolist():
                data = source_zip.read(item.filename)
                if item.filename == "word/document.xml":
                    data = mutated_document_xml.encode("utf-8")
                target_zip.writestr(item, data)

        restored = ReportExportService._restore_template_header_footer(
            buffer.getvalue(),
            ReportExportService.WAIS3_TEMPLATE_PATH,
        )

        with ZipFile(BytesIO(restored), "r") as archive:
            restored_document_xml = archive.read("word/document.xml").decode("utf-8")

        self.assertIn("sectPr", restored_document_xml)
        self.assertIn("headerReference", restored_document_xml)
        self.assertIn("footerReference", restored_document_xml)

    def test_rebuild_qualitative_section_preserves_template_chart_for_wais3(self):
        template = Document(str(ReportExportService.WAIS3_TEMPLATE_PATH))
        template_chart_blocks = ReportExportService._extract_template_chart_blocks(template)

        document = Document()
        start = document.add_paragraph("ANÁLISE QUALITATIVA")
        start._p.addnext(deepcopy(template_chart_blocks[0]))
        end = document.add_paragraph("Conclusão")
        start._p.getnext().addnext(end._p)

        ReportExportService._rebuild_qualitative_section(
            document,
            sections={},
            context={
                "patient": {"sex": "F"},
                "validated_tests": [
                    {
                        "instrument_code": "wais3",
                        "structured_results": {
                            "indices": {
                                "qi_total": {
                                    "pontuacao_composta": 95,
                                    "classificacao": "Média",
                                }
                            }
                        },
                    }
                ],
            },
        )

        self.assertTrue(ReportExportService._extract_template_chart_blocks(document))

    def test_rebuild_qualitative_section_uses_wais3_model_titles_and_order(self):
        template = Document(str(ReportExportService.WAIS3_TEMPLATE_PATH))
        template_chart_blocks = ReportExportService._extract_template_chart_blocks(template)

        document = Document()
        start = document.add_paragraph("ANÁLISE QUALITATIVA")
        start._p.addnext(deepcopy(template_chart_blocks[0]))
        end = document.add_paragraph("Conclusão")
        start._p.getnext().addnext(end._p)

        ReportExportService._rebuild_qualitative_section(
            document,
            sections={
                "eficiencia_intelectual": "Interpretação global do WAIS3.",
                "linguagem": "Interpretação da linguagem.",
            },
            context={
                "patient": {"sex": "M"},
                "validated_tests": [
                    {
                        "instrument_code": "wais3",
                        "structured_results": {
                            "indices": {
                                "qi_total": {
                                    "pontuacao_composta": 95,
                                    "classificacao": "Média",
                                },
                                "qi_verbal": {
                                    "pontuacao_composta": 93,
                                    "classificacao": "Média",
                                },
                                "qi_execucao": {
                                    "pontuacao_composta": 91,
                                    "classificacao": "Média",
                                },
                                "compreensao_verbal": {
                                    "pontuacao_composta": 94,
                                    "classificacao": "Média",
                                },
                                "organizacao_perceptual": {
                                    "pontuacao_composta": 92,
                                    "classificacao": "Média",
                                },
                                "memoria_operacional": {
                                    "pontuacao_composta": 96,
                                    "classificacao": "Média",
                                },
                                "velocidade_processamento": {
                                    "pontuacao_composta": 90,
                                    "classificacao": "Média",
                                },
                            }
                        },
                        "wais3_tables": {
                            "linguagem": [
                                {
                                    "label": "Semelhanças",
                                    "maxScore": "38",
                                    "avgScore": "17-28",
                                    "minScore": "9-10",
                                    "obtainedScore": "18",
                                    "classification": "Média",
                                }
                            ]
                        },
                    }
                ],
            },
        )

        texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        self.assertIn("Desempenho do paciente no WAIS III", texts)
        self.assertIn("Subescalas WAIS III", texts)
        self.assertIn("Gráfico 1 WAIS III - INDICES DE QIS", texts)
        self.assertNotIn("5.1. Desempenho do paciente no WAIS-III", texts)

        desempenho_index = texts.index("Desempenho do paciente no WAIS III")
        interpretacao_index = next(
            i for i, text in enumerate(texts) if "Interpretação global do WAIS3." in text
        )
        subescalas_index = texts.index("Subescalas WAIS III")

        self.assertLess(desempenho_index, interpretacao_index)
        self.assertLess(interpretacao_index, subescalas_index)

    def test_rebuild_qualitative_section_uses_wais3_skill_labels_for_late_chapters(self):
        template = Document(str(ReportExportService.WAIS3_TEMPLATE_PATH))
        template_chart_blocks = ReportExportService._extract_template_chart_blocks(template)

        document = Document()
        start = document.add_paragraph("ANÁLISE QUALITATIVA")
        start._p.addnext(deepcopy(template_chart_blocks[0]))
        end = document.add_paragraph("Conclusão")
        start._p.getnext().addnext(end._p)

        ReportExportService._rebuild_qualitative_section(
            document,
            sections={
                "etdah_ad": "Interpretacao do ETDAH-AD.",
                "srs2": "Interpretacao do SRS-2.",
                "aspectos_emocionais_comportamentais": "Interpretacao da EBADEP-A.",
            },
            context={
                "patient": {"sex": "F"},
                "validated_tests": [
                    {"instrument_code": "wais3", "structured_results": {"indices": {"qi_total": {"pontuacao_composta": 95, "classificacao": "Média"}}}},
                    {"instrument_code": "etdah_ad"},
                    {"instrument_code": "ebadep_a"},
                    {"instrument_code": "srs2"},
                ],
            },
        )

        texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        self.assertIn("Função Executiva", texts)
        self.assertTrue(any(text.startswith("A E-TDAH-AD e um instrumento de autorrelato destinado a adolescentes e adultos") for text in texts))
        self.assertTrue(any(text.startswith("A EBADEP-A avalia a presenca e a intensidade de sintomas depressivos em adultos") for text in texts))
        self.assertTrue(any(text.startswith("A SRS-2 e uma escala destinada a investigacao de dificuldades") for text in texts))

    def test_rebuild_qualitative_section_does_not_duplicate_wais3_global_intro(self):
        document = Document()
        start = document.add_paragraph("ANÁLISE QUALITATIVA")
        end = document.add_paragraph("Conclusão")
        start._p.addnext(end._p)

        ReportExportService._rebuild_qualitative_section(
            document,
            sections={},
            context={
                "patient": {"sex": "M"},
                "validated_tests": [
                    {
                        "instrument_code": "wais3",
                        "structured_results": {
                            "indices": {
                                "qi_total": {
                                    "pontuacao_composta": 98,
                                    "classificacao": "Média",
                                },
                                "qi_verbal": {
                                    "pontuacao_composta": 99,
                                    "classificacao": "Média",
                                },
                                "qi_execucao": {
                                    "pontuacao_composta": 99,
                                    "classificacao": "Média",
                                },
                                "compreensao_verbal": {
                                    "pontuacao_composta": 102,
                                    "classificacao": "Média",
                                },
                                "organizacao_perceptual": {
                                    "pontuacao_composta": 96,
                                    "classificacao": "Média",
                                },
                                "memoria_operacional": {
                                    "pontuacao_composta": 90,
                                    "classificacao": "Média",
                                },
                                "velocidade_processamento": {
                                    "pontuacao_composta": 105,
                                    "classificacao": "Média",
                                },
                            }
                        },
                    }
                ],
            },
        )

        global_lines = [
            p.text.strip()
            for p in document.paragraphs
            if p.text.strip().startswith("Capacidade Cognitiva Global:")
        ]

        self.assertEqual(len(global_lines), 1)
        self.assertEqual(global_lines[0].count("Capacidade Cognitiva Global:"), 1)

    def test_rebuild_qualitative_section_uses_wisc4_model_headings(self):
        document = Document()
        start = document.add_paragraph("ANÁLISE QUALITATIVA")
        end = document.add_paragraph("Conclusão")
        start._p.addnext(end._p)

        ReportExportService._rebuild_qualitative_section(
            document,
            sections={},
            context={
                "patient": {"full_name": "Maria Clara", "sex": "F", "birth_date": "2016-01-10"},
                "validated_tests": [
                    {
                        "instrument_code": "wisc4",
                        "structured_results": {
                            "qit_data": {"escore_composto": 61, "classificacao": "Extremamente Baixo"},
                            "indices": [
                                {"indice": "icv", "escore_composto": 84, "classificacao": "Média Inferior"},
                                {"indice": "iop", "escore_composto": 61, "classificacao": "Extremamente Baixo"},
                                {"indice": "imt", "escore_composto": 52, "classificacao": "Extremamente Baixo"},
                                {"indice": "ivp", "escore_composto": 68, "classificacao": "Extremamente Baixo"},
                            ],
                            "subtestes": [
                                {"codigo": "SM", "subteste": "Semelhanças", "escore_bruto": 23, "classificacao": "Média"},
                                {"codigo": "CN", "subteste": "Conceitos Figurativos", "escore_bruto": 11, "classificacao": "Média Inferior"},
                                {"codigo": "CO", "subteste": "Compreensão", "escore_bruto": 24, "classificacao": "Média"},
                                {"codigo": "RM", "subteste": "Raciocínio Matricial", "escore_bruto": 12, "classificacao": "Limítrofe"},
                                {"codigo": "VC", "subteste": "Vocabulário", "escore_bruto": 10, "classificacao": "Limítrofe"},
                                {"codigo": "CB", "subteste": "Cubos", "escore_bruto": 8, "classificacao": "Extremamente Baixo"},
                                {"codigo": "SNL", "subteste": "Sequência de Números e Letras", "escore_bruto": 5, "classificacao": "Extremamente Baixo"},
                                {"codigo": "DG", "subteste": "Dígitos", "escore_bruto": 7, "classificacao": "Limítrofe"},
                            ],
                        },
                        "classified_payload": {
                            "subtestes": [
                                {"codigo": "SM", "subteste": "Semelhanças", "escore_bruto": 23, "classificacao": "Média"},
                                {"codigo": "CN", "subteste": "Conceitos Figurativos", "escore_bruto": 11, "classificacao": "Média Inferior"},
                                {"codigo": "CO", "subteste": "Compreensão", "escore_bruto": 24, "classificacao": "Média"},
                                {"codigo": "RM", "subteste": "Raciocínio Matricial", "escore_bruto": 12, "classificacao": "Limítrofe"},
                                {"codigo": "VC", "subteste": "Vocabulário", "escore_bruto": 10, "classificacao": "Limítrofe"},
                                {"codigo": "CB", "subteste": "Cubos", "escore_bruto": 8, "classificacao": "Extremamente Baixo"},
                                {"codigo": "SNL", "subteste": "Sequência de Números e Letras", "escore_bruto": 5, "classificacao": "Extremamente Baixo"},
                                {"codigo": "DG", "subteste": "Dígitos", "escore_bruto": 7, "classificacao": "Limítrofe"},
                            ]
                        },
                        "applied_on": "2024-04-20",
                    }
                ],
                "evaluation": {"start_date": "2024-04-20"},
            },
        )

        texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        self.assertIn("Capacidade Cognitiva Global", texts)
        self.assertIn("Desempenho da paciente no WISC-IV", texts)
        self.assertIn("Subescalas WISC-IV", texts)
        self.assertIn("Função Executiva", texts)
        self.assertTrue(any(text.startswith("Interpretação e Observações Clínicas:") for text in texts))

    def test_build_adolescent_document_uses_model_wisc4_titles_when_optional_tests_are_missing(self):
        document = ReportExportService._build_adolescent_document(
            self._report_stub(),
            sections={"conclusao": "Conclusão clínica.", "sugestoes_conduta": "- Sugestão 1"},
            context={
                "patient": {"full_name": "Maria Clara", "sex": "F", "birth_date": "2016-01-10"},
                "evaluation": {"start_date": "2024-04-20"},
                "validated_tests": [
                    {
                        "instrument_code": "wisc4",
                        "structured_results": {
                            "qit_data": {"escore_composto": 61, "classificacao": "Extremamente Baixo"},
                            "indices": [
                                {"indice": "icv", "escore_composto": 84, "classificacao": "Média Inferior"},
                                {"indice": "iop", "escore_composto": 61, "classificacao": "Extremamente Baixo"},
                                {"indice": "imt", "escore_composto": 52, "classificacao": "Extremamente Baixo"},
                                {"indice": "ivp", "escore_composto": 68, "classificacao": "Extremamente Baixo"},
                            ],
                            "subtestes": [
                                {"codigo": "SM", "subteste": "Semelhanças", "escore_bruto": 23, "classificacao": "Média"},
                                {"codigo": "CN", "subteste": "Conceitos Figurativos", "escore_bruto": 11, "classificacao": "Média Inferior"},
                                {"codigo": "CO", "subteste": "Compreensão", "escore_bruto": 24, "classificacao": "Média"},
                                {"codigo": "RM", "subteste": "Raciocínio Matricial", "escore_bruto": 12, "classificacao": "Limítrofe"},
                                {"codigo": "VC", "subteste": "Vocabulário", "escore_bruto": 10, "classificacao": "Limítrofe"},
                                {"codigo": "CB", "subteste": "Cubos", "escore_bruto": 8, "classificacao": "Extremamente Baixo"},
                                {"codigo": "SNL", "subteste": "Sequência de Números e Letras", "escore_bruto": 5, "classificacao": "Extremamente Baixo"},
                                {"codigo": "DG", "subteste": "Dígitos", "escore_bruto": 7, "classificacao": "Limítrofe"},
                            ],
                        },
                        "classified_payload": {
                            "subtestes": [
                                {"codigo": "SM", "subteste": "Semelhanças", "escore_bruto": 23, "classificacao": "Média"}
                            ]
                        },
                        "applied_on": "2024-04-20",
                    }
                ],
            },
        )

        texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        self.assertIn("5. ANÁLISE QUALITATIVA", texts)
        self.assertIn("Capacidade Cognitiva Global", texts)
        self.assertIn("Desempenho da paciente no WISC-IV", texts)
        self.assertIn("Subescalas WISC-IV", texts)
        self.assertIn("Conclusão", texts)
        self.assertIn("Sugestões de Conduta (Encaminhamentos)", texts)
        self.assertIn("Referencia Bibliográfica", texts)
        self.assertNotIn("6. BPA-2 – BATERIA PSICOLÓGICA PARA AVALIAÇÃO DA ATENÇÃO", texts)
        self.assertNotIn("6. CONCLUSÃO", texts)
        self.assertNotIn("7. SUGESTÕES DE CONDUTA (ENCAMINHAMENTOS)", texts)
        self.assertNotIn("8. REFERÊNCIA BIBLIOGRÁFICA", texts)

    def test_replace_simple_sections_uses_wais3_skill_structure_for_demand_and_procedures(self):
        document = Document()
        document.add_paragraph("IDENTIFICAÇÃO")
        document.add_paragraph("DESCRIÇÃO DA DEMANDA")
        document.add_paragraph("PROCEDIMENTOS")
        document.add_paragraph("ANÁLISE")

        ReportExportService._replace_simple_sections(
            document,
            self._report_stub(),
            sections={
                "identificacao": "Identificação pronta.",
                "descricao_demanda": "O paciente Marcos foi encaminhado para avaliação neuropsicológica por apresentar queixas relacionadas a desatenção e esquecimentos frequentes.",
                "procedimentos": "Para esta avaliação foram realizadas: uma sessão de anamnese, 05 sessões de testagem com o paciente e uma sessão de devolutiva.",
                "historia_pessoal": "História pessoal resumida.",
            },
            context={
                "patient": {"full_name": "Marcos Henrique Carvalho Santos"},
                "validated_tests": [
                    {"instrument_code": "wais3"},
                    {"instrument_code": "bpa2"},
                ],
            },
        )

        texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        self.assertIn("PROCEDIMENTOS", texts)

    def test_replace_simple_sections_uses_default_wais3_demand_text_when_empty(self):
        document = Document()
        document.add_paragraph("IDENTIFICAÇÃO")
        document.add_paragraph("DESCRIÇÃO DA DEMANDA")
        document.add_paragraph("PROCEDIMENTOS")
        document.add_paragraph("ANÁLISE")

        ReportExportService._replace_simple_sections(
            document,
            self._report_stub(),
            sections={},
            context={"patient": {"full_name": "Marcos Henrique Carvalho Santos"}, "validated_tests": []},
        )

        texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        self.assertIn("Motivo do Encaminhamento", texts)
        self.assertIn(
            "Levando em consideração a demanda apresentada, objetivou-se avaliar as funções cognitivas, aspectos comportamentais, emocionais e sociais do paciente, a fim de compreender seu funcionamento global e subsidiar orientações para o manejo adequado em contextos familiar e escolar.",
            texts,
        )

    def test_replace_simple_sections_rebuilds_identification_chapter_with_fixed_labels(self):
        document = Document()
        document.add_paragraph("IDENTIFICAÇÃO")
        document.add_paragraph("DESCRIÇÃO DA DEMANDA")

        report = self._report_stub()
        report.interested_party = "Dra. Exemplo"
        report.purpose = "Auxílio diagnóstico"

        ReportExportService._replace_simple_sections(
            document,
            report,
            sections={},
            context={
                "patient": {
                    "full_name": "Marcos Henrique Carvalho Santos",
                    "sex": "M",
                    "birth_date": "2002-09-12",
                    "schooling": "superior incompleto",
                }
            },
        )

        texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]

        self.assertEqual(texts[1], "1.1. Identificação do laudo:")
        self.assertEqual(texts[2], "Autora: Jacqueline Oliveira Caires (CRP 09/6017)")
        self.assertEqual(texts[3], "Interessado: Dra. Exemplo")
        self.assertEqual(texts[4], "Finalidade: Auxílio diagnóstico")
        self.assertEqual(texts[5], "1.2. Identificação do paciente:")
        self.assertEqual(texts[6], "Nome: Marcos Henrique Carvalho Santos")
        self.assertEqual(texts[7], "Sexo: Masculino")
        self.assertEqual(texts[8], "Data de nascimento: 12/09/2002")
        self.assertTrue(texts[9].startswith("Idade: "))
        self.assertIn("Filiação: Não informada", texts)
        self.assertIn("Escolaridade: superior incompleto", texts)

    def test_sanitize_chart_xml_inlines_cached_refs_and_removes_external_data(self):
        chart_xml = b'''<?xml version="1.0" encoding="UTF-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <c:externalData r:id="rId1"/>
  <c:chart>
    <c:plotArea>
      <c:barChart>
        <c:ser>
          <c:idx val="0"/>
          <c:order val="0"/>
          <c:tx>
            <c:strRef>
              <c:f>Plan1!$B$1</c:f>
              <c:strCache><c:ptCount val="1"/><c:pt idx="0"><c:v>Serie A</c:v></c:pt></c:strCache>
            </c:strRef>
          </c:tx>
          <c:cat>
            <c:strRef>
              <c:f>Plan1!$A$2:$A$3</c:f>
              <c:strCache><c:ptCount val="2"/><c:pt idx="0"><c:v>A</c:v></c:pt><c:pt idx="1"><c:v>B</c:v></c:pt></c:strCache>
            </c:strRef>
          </c:cat>
          <c:val>
            <c:numRef>
              <c:f>Plan1!$B$2:$B$3</c:f>
              <c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="2"/><c:pt idx="0"><c:v>10</c:v></c:pt><c:pt idx="1"><c:v>20</c:v></c:pt></c:numCache>
            </c:numRef>
          </c:val>
        </c:ser>
      </c:barChart>
    </c:plotArea>
  </c:chart>
</c:chartSpace>'''

        sanitized = ReportExportService._sanitize_chart_xml_bytes(chart_xml)
        root = ET.fromstring(sanitized)
        ns = ReportExportService.CHART_NS

        self.assertIsNone(root.find('.//c:externalData', ns))
        self.assertIsNone(root.find('.//c:strRef', ns))
        self.assertIsNone(root.find('.//c:numRef', ns))
        self.assertEqual(root.find('.//c:tx/c:v', ns).text, 'Serie A')
        self.assertEqual(
            [node.text for node in root.findall('.//c:cat/c:strLit/c:pt/c:v', ns)],
            ['A', 'B'],
        )
        self.assertEqual(
            [node.text for node in root.findall('.//c:val/c:numLit/c:pt/c:v', ns)],
            ['10', '20'],
        )

    def test_strip_external_chart_relationships_keeps_internal_links(self):
        rels_xml = b'''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.microsoft.com/office/2011/relationships/chartStyle" Target="style1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject" Target="https://example.com/test.xlsx" TargetMode="External"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartUserShapes" Target="../drawings/drawing1.xml"/>
</Relationships>'''

        sanitized = ReportExportService._strip_external_chart_relationships(rels_xml)
        text = sanitized.decode('utf-8')

        self.assertIn('style1.xml', text)
        self.assertIn('../drawings/drawing1.xml', text)
        self.assertNotIn('example.com/test.xlsx', text)
        self.assertNotIn('oleObject', text)

    def test_sanitize_chart_xml_flattens_multi_level_categories(self):
        chart_xml = b'''<?xml version="1.0" encoding="UTF-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">
  <c:chart>
    <c:plotArea>
      <c:barChart>
        <c:ser>
          <c:idx val="0"/>
          <c:order val="0"/>
          <c:cat>
            <c:multiLvlStrRef>
              <c:f>Plan1!$A$1:$B$3</c:f>
              <c:multiLvlStrCache>
                <c:ptCount val="3"/>
                <c:lvl>
                  <c:pt idx="0"><c:v>A</c:v></c:pt>
                  <c:pt idx="1"><c:v>B</c:v></c:pt>
                </c:lvl>
                <c:lvl>
                  <c:pt idx="2"><c:v>C</c:v></c:pt>
                </c:lvl>
              </c:multiLvlStrCache>
            </c:multiLvlStrRef>
          </c:cat>
          <c:val>
            <c:numRef>
              <c:f>Plan1!$C$1:$C$3</c:f>
              <c:numCache><c:ptCount val="3"/><c:pt idx="0"><c:v>1</c:v></c:pt><c:pt idx="1"><c:v>2</c:v></c:pt><c:pt idx="2"><c:v>3</c:v></c:pt></c:numCache>
            </c:numRef>
          </c:val>
        </c:ser>
      </c:barChart>
    </c:plotArea>
  </c:chart>
</c:chartSpace>'''

        sanitized = ReportExportService._sanitize_chart_xml_bytes(chart_xml)
        root = ET.fromstring(sanitized)
        ns = ReportExportService.CHART_NS

        self.assertIsNone(root.find('.//c:multiLvlStrRef', ns))
        self.assertEqual(
            [node.text for node in root.findall('.//c:cat/c:strLit/c:pt/c:v', ns)],
            ['A', 'B', 'C'],
        )

    def test_sanitize_chart_xml_preserves_mc_requires_namespace_prefix(self):
        chart_xml = b'''<?xml version="1.0" encoding="UTF-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:c16r2="http://schemas.microsoft.com/office/drawing/2015/06/chart">
  <mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">
    <mc:Choice Requires="c14" xmlns:c14="http://schemas.microsoft.com/office/drawing/2007/8/2/chart">
      <c14:style val="102"/>
    </mc:Choice>
    <mc:Fallback><c:style val="2"/></mc:Fallback>
  </mc:AlternateContent>
  <c:chart><c:plotArea/></c:chart>
</c:chartSpace>'''

        sanitized = ReportExportService._sanitize_chart_xml_bytes(chart_xml).decode('utf-8')

        self.assertIn('Requires="c14"', sanitized)
        self.assertIn('xmlns:c14="http://schemas.microsoft.com/office/drawing/2007/8/2/chart"', sanitized)
        self.assertIn('<c14:style', sanitized)

    def test_update_chart_series_before_sanitize_replaces_template_values(self):
        chart_xml = LET.fromstring(b'''<?xml version="1.0" encoding="UTF-8"?>
<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart">
  <c:chart>
    <c:plotArea>
      <c:barChart>
        <c:ser>
          <c:idx val="0"/>
          <c:order val="0"/>
          <c:cat>
            <c:strRef>
              <c:f>Plan1!$A$1:$A$2</c:f>
              <c:strCache>
                <c:ptCount val="2"/>
                <c:pt idx="0"><c:v>Old A</c:v></c:pt>
                <c:pt idx="1"><c:v>Old B</c:v></c:pt>
              </c:strCache>
            </c:strRef>
          </c:cat>
          <c:val>
            <c:numRef>
              <c:f>Plan1!$B$1:$B$2</c:f>
              <c:numCache>
                <c:formatCode>General</c:formatCode>
                <c:ptCount val="2"/>
                <c:pt idx="0"><c:v>84</c:v></c:pt>
                <c:pt idx="1"><c:v>61</c:v></c:pt>
              </c:numCache>
            </c:numRef>
          </c:val>
        </c:ser>
      </c:barChart>
    </c:plotArea>
  </c:chart>
</c:chartSpace>''')

        ReportExportService._update_chart_series(chart_xml, 0, ['ICV', 'IOP'], [113, 90])
        ReportExportService._sanitize_chart_root(chart_xml)

        ns = ReportExportService.CHART_NS
        self.assertEqual(
            [node.text for node in chart_xml.findall('.//c:cat//c:pt/c:v', ns)],
            ['ICV', 'IOP'],
        )
        self.assertEqual(
            [node.text for node in chart_xml.findall('.//c:val//c:pt/c:v', ns)],
            ['113', '90'],
        )

    def test_document_chart_targets_follow_document_order(self):
        buffer = BytesIO()
        with ZipFile(buffer, 'w') as docx:
            docx.writestr(
                'word/document.xml',
                '''<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <w:body>
    <w:p><w:r><w:drawing><c:chart r:id="rId9"/></w:drawing></w:r></w:p>
    <w:p><w:r><w:drawing><c:chart r:id="rId3"/></w:drawing></w:r></w:p>
    <w:p><w:r><w:drawing><c:chart r:id="rId11"/></w:drawing></w:r></w:p>
  </w:body>
</w:document>''',
            )
            docx.writestr(
                'word/_rels/document.xml.rels',
                '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart7.xml"/>
  <Relationship Id="rId9" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart2.xml"/>
  <Relationship Id="rId11" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart10.xml"/>
</Relationships>''',
            )

        self.assertEqual(
            ReportExportService._document_chart_targets(buffer.getvalue()),
            [
                'word/charts/chart2.xml',
                'word/charts/chart7.xml',
                'word/charts/chart10.xml',
            ],
        )

    def test_fdt_chart_payload_uses_table_order_not_saved_chart_order(self):
        test = {
            'classified_payload': {
                'metric_results': [
                    {'codigo': 'contagem', 'media': 68.91, 'valor': 30, 'percentil_num': 12},
                    {'codigo': 'leitura', 'media': 37.44, 'valor': 23.3, 'percentil_num': 18},
                    {'codigo': 'escolha', 'media': 127, 'valor': 47.8, 'percentil_num': 20},
                    {'codigo': 'alternancia', 'media': 235, 'valor': 56.9, 'percentil_num': 22},
                    {'codigo': 'inibicao', 'media': 89.56, 'valor': 23.8, 'percentil_num': 24},
                    {'codigo': 'flexibilidade', 'media': 197.56, 'valor': 33.6, 'percentil_num': 26},
                ],
                'erros': {
                    'contagem': {'qtde_erros': 1},
                    'leitura': {'qtde_erros': 0},
                    'escolha': {'qtde_erros': 10},
                    'alternancia': {'qtde_erros': 16},
                },
                'charts': {
                    'automaticos': {
                        'categories': ['saved'],
                        'series': [{'values': [999]}],
                    },
                },
            }
        }

        automatic_categories, automatic_series = ReportExportService._fdt_chart_payload(test, automatic=True)
        controlled_categories, controlled_series = ReportExportService._fdt_chart_payload(test, automatic=False)

        self.assertEqual(
            automatic_categories,
            ['Tempo Médio', 'Tempo Obtido', 'Erros', 'Desempenho %', 'Indicativo de Déficit', 'Indicativo de Dificuldade Discreta', 'Sem Indicativo de Déficit'],
        )
        self.assertEqual(automatic_series[0], [68.91, 30.0, 1.0, 5.0, 5.0, 25.0, 50.0])
        self.assertEqual(automatic_series[1], [37.44, 23.3, 0.0, 5.0, 5.0, 25.0, 50.0])
        self.assertEqual(controlled_series[0], [127.0, 47.8, 10.0, 5.0, 5.0, 25.0, 50.0])
        self.assertEqual(controlled_series[3], [197.56, 33.6, 0.0, 25.0, 5.0, 25.0, 50.0])

    def test_ravlt_chart_payload_accepts_pt_br_numeric_strings(self):
        service = ReportExportService
        original = service._ravlt_rows
        service._ravlt_rows = classmethod(lambda cls, test, context=None: [
            ['Desempenho', 'A1', 'RET', 'I.P.'],
            ['Esperado', '6', '1', '0,88'],
            ['Mínimo', '5', '0,89', '0,73'],
            ['Obtido', '6', '0,75', '0,86'],
        ])
        try:
            categories, series = service._ravlt_chart_payload({}, None)
        finally:
            service._ravlt_rows = original

        self.assertEqual(categories, ['A1', 'RET', 'I.P.'])
        self.assertEqual(series, [[6.0, 1.0, 0.88], [5.0, 0.89, 0.73], [6.0, 0.75, 0.86]])

    def test_prune_unused_chart_parts_removes_unreferenced_chart_files(self):
        buffer = BytesIO()
        with ZipFile(buffer, 'w') as docx:
            docx.writestr(
                '[Content_Types].xml',
                '''<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/charts/chart1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>
  <Override PartName="/word/charts/chart2.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>
</Types>''',
            )
            docx.writestr(
                'word/document.xml',
                '''<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:body><w:p><w:r><w:drawing><c:chart r:id="rId1"/></w:drawing></w:r></w:p></w:body></w:document>''',
            )
            docx.writestr(
                'word/_rels/document.xml.rels',
                '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/>
</Relationships>''',
            )
            docx.writestr('word/charts/chart1.xml', '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>')
            docx.writestr('word/charts/chart2.xml', '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>')
            docx.writestr('word/charts/_rels/chart1.xml.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')
            docx.writestr('word/charts/_rels/chart2.xml.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>')

        pruned = ReportExportService._prune_unused_chart_parts(buffer.getvalue())
        with ZipFile(BytesIO(pruned), 'r') as docx:
            names = set(docx.namelist())
            self.assertIn('word/charts/chart1.xml', names)
            self.assertNotIn('word/charts/chart2.xml', names)
            self.assertNotIn('word/charts/_rels/chart2.xml.rels', names)
            self.assertNotIn('/word/charts/chart2.xml', docx.read('[Content_Types].xml').decode('utf-8'))
            rels_text = docx.read('word/_rels/document.xml.rels').decode('utf-8')
            self.assertIn('charts/chart1.xml', rels_text)
            self.assertNotIn('charts/chart2.xml', rels_text)

    def test_prune_unused_chart_parts_keeps_chart_user_shapes_targets(self):
        buffer = BytesIO()
        with ZipFile(buffer, 'w') as docx:
            docx.writestr(
                '[Content_Types].xml',
                '''<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
  <Override PartName="/word/charts/chart1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>
  <Override PartName="/word/drawings/drawing1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawing+xml"/>
</Types>''',
            )
            docx.writestr(
                'word/document.xml',
                '''<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><w:body><w:p><w:r><w:drawing><c:chart r:id="rId1"/></w:drawing></w:r></w:p></w:body></w:document>''',
            )
            docx.writestr(
                'word/_rels/document.xml.rels',
                '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/>
</Relationships>''',
            )
            docx.writestr('word/charts/chart1.xml', '<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"/>')
            docx.writestr(
                'word/charts/_rels/chart1.xml.rels',
                '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chartUserShapes" Target="../drawings/drawing1.xml"/>
</Relationships>''',
            )
            docx.writestr('word/drawings/drawing1.xml', '<xdr:wsDr xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"/>')

        pruned = ReportExportService._prune_unused_chart_parts(buffer.getvalue())
        with ZipFile(BytesIO(pruned), 'r') as docx:
            names = set(docx.namelist())
            self.assertIn('word/drawings/drawing1.xml', names)

    def test_append_body_element_before_sectpr_keeps_section_properties_last(self):
        document = Document()
        document.add_paragraph('Antes')

        new_paragraph = parse_xml(
            '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:r><w:t>Chart placeholder</w:t></w:r></w:p>'
        )

        ReportExportService._append_body_element_before_sectpr(document, new_paragraph)

        body_children = list(document._body._element.iterchildren())
        self.assertEqual(body_children[-1].tag.split('}')[-1], 'sectPr')
        self.assertEqual(body_children[-2].tag.split('}')[-1], 'p')
