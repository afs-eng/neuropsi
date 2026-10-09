from io import BytesIO
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from django.test import SimpleTestCase
from docx import Document

from apps.reports.builders.srs2_profiles import comparable_profiles
from apps.reports.services.report_export_service import ReportExportService
from apps.reports.services.section_context_service import SectionContextService
from apps.reports.services.patient_identity_service import PatientIdentityService
from apps.reports.services.wais3_standardization import WAIS3StandardizationService
from apps.reports import tests as report_tests
from apps.tests.srs2.interpreters import DISPLAY_NAMES, INTERPRETATION_ORDER


class WAIS3ApprovedModelTests(SimpleTestCase):
    def context(self, primary="wais3", age=32):
        return report_tests.ReportExportChartSanitizationTests._export_context(self, primary, age)

    def export(self, context, edited=None):
        sections = [SimpleNamespace(key=key, content_edited=value, content_generated="") for key, value in (edited or {}).items()]
        report = SimpleNamespace(sections=SimpleNamespace(all=lambda: sections), patient=None)
        with patch("apps.reports.services.report_export_service.ReportContextService.sync_report_context", return_value=context):
            return ReportExportService.generate_docx_bytes(report)

    def srs2_test(self, form, respondent, scores):
        return {"instrument_code": "srs2", "raw_payload": {"form": form, "respondent_name": respondent}, "classified_payload": {
            "form": form,
            "resultados": [{"variável": key, "nome": DISPLAY_NAMES[key], "tscore": score} for key, score in zip(INTERPRETATION_ORDER, scores)],
        }}

    def test_patient_free_template_contains_only_layout_and_empty_chart_data(self):
        with ZipFile(ReportExportService.WAIS3_STANDARD_TEMPLATE_PATH) as package:
            self.assertFalse(any(name.startswith("word/embeddings/") for name in package.namelist()))
            for name in package.namelist():
                if name.endswith(".xml"):
                    self.assertNotIn(b"mascarenhas", package.read(name).lower())
                    self.assertNotIn(b"mariana", package.read(name).lower())
                if name.startswith("word/charts/chart") and name.endswith(".xml"):
                    root = ET.fromstring(package.read(name))
                    values = root.findall(".//c:val//c:v", ReportExportService.CHART_NS)
                    self.assertTrue(all(value.text == "0" for value in values))
        document = Document(ReportExportService.WAIS3_STANDARD_TEMPLATE_PATH)
        self.assertFalse(any(paragraph.text.strip() for paragraph in document.paragraphs))

    def test_model_has_approved_margins_headings_subscale_order_and_closing(self):
        document = Document(BytesIO(self.export(self.context())))
        section = document.sections[0]
        self.assertEqual([round(value.cm, 1) for value in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin)], [3, 2.5, 2, 2])
        headings = [paragraph.text for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
        self.assertEqual(headings, ["IDENTIFICAÇÃO", "DESCRIÇÃO DA DEMANDA", "PROCEDIMENTOS", "ANÁLISE", "ANÁLISE QUALITATIVA", "CONCLUSÃO", "SUGESTÕES DE CONDUTA (ENCAMINHAMENTOS):", "CONSIDERAÇÕES FINAIS", "REFERÊNCIAS BIBLIOGRÁFICAS"])
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertLess(texts.index("Linguagem"), texts.index("Gnosias e Praxias"))
        self.assertLess(texts.index("Gnosias e Praxias"), texts.index("Função Executiva"))
        self.assertLess(texts.index("Função Executiva"), texts.index("Memória e Aprendizagem"))
        self.assertIn("Interessado: O paciente", texts)
        self.assertIn("Neuropsicóloga - CRP 09/6017", texts)
        self.assertNotIn("Eric", "\n".join(texts))
        self.assertTrue(any(section.header._element.xpath(".//w:drawing") for section in document.sections))

    def test_revised_clinical_texts_are_preserved(self):
        edited = {key: f"Texto revisado pelo profissional número {index}." for index, key in enumerate(("capacidade_cognitiva_global", "linguagem", "conclusao", "bpa2", "sugestoes_conduta"), 1)}
        document = Document(BytesIO(self.export(self.context(), edited)))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        for content in edited.values():
            self.assertIn(content, text)

    def test_autorrelato_heterorrelato_and_comparison_use_each_saved_profile(self):
        context = self.context()
        context["validated_tests"] = [test for test in context["validated_tests"] if test["instrument_code"] != "srs2"]
        expected = [[55, 66, 72, 68, 69, 70, 70], [62, 60, 59, 64, 63, 62, 62]]
        context["validated_tests"].extend([self.srs2_test("adulto_autorrelato", "Respondente A", expected[0]), self.srs2_test("adulto_heterorrelato", "Respondente B", expected[1])])
        data = self.export(context, {"srs2": "Análise conjunta revisada pelo profissional."})
        document = Document(BytesIO(data))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Autorrelato — Respondente A", text)
        self.assertIn("Heterorrelato — Respondente B", text)
        self.assertIn("Análise conjunta revisada pelo profissional.", text)
        targets = ReportExportService._document_chart_targets(data)
        self.assertEqual(len(targets), 9)
        with ZipFile(BytesIO(data)) as package:
            for target, values in zip(targets[-3:-1], expected):
                root = ET.fromstring(package.read(target))
                self.assertEqual([float(point.text) for point in root.findall(".//c:val//c:pt/c:v", ReportExportService.CHART_NS)], values)
            comparison = ET.fromstring(package.read(targets[-1]))
            series = comparison.findall(".//c:ser", ReportExportService.CHART_NS)
            self.assertEqual(len(series), 2)
            for node, values in zip(series, expected):
                self.assertEqual([float(point.text) for point in node.findall("c:val//c:pt/c:v", ReportExportService.CHART_NS)], values)

    def test_comparison_is_omitted_for_incompatible_or_unidentified_forms(self):
        tests = [self.srs2_test("adulto_autorrelato", "A", [60] * 7), self.srs2_test("idade_escolar", "B", [60] * 7)]
        self.assertEqual(comparable_profiles(tests), ([], [], []))
        tests[1]["raw_payload"]["form"] = None
        tests[1]["classified_payload"]["form"] = None
        self.assertEqual(comparable_profiles(tests), ([], [], []))

    def test_comparison_aligns_domains_by_identity_not_list_position(self):
        first = self.srs2_test("adulto_autorrelato", "A", [55, 66, 72, 68, 69, 70, 70])
        second = self.srs2_test("adulto_heterorrelato", "B", [62, 60, 59, 64, 63, 62, 62])
        second["classified_payload"]["resultados"].reverse()
        second["classified_payload"]["resultados"].pop()
        categories, values, _ = comparable_profiles([first, second])
        self.assertNotIn(DISPLAY_NAMES["percepção_social"], categories)
        self.assertEqual(values[1], [60, 59, 64, 63, 62, 62])

    def test_registered_respondent_name_is_not_replaced_with_the_patient_name(self):
        context = self.context()
        context["validated_tests"] = [test for test in context["validated_tests"] if test["instrument_code"] != "srs2"]
        context["validated_tests"].append(self.srs2_test("adulto_heterorrelato", "Maria Silva", [62] * 7))
        document = Document(BytesIO(self.export(context, {"srs2": "O heterorrelato foi respondido por Maria Silva."})))
        text = "\n".join(paragraph.text for paragraph in document.paragraphs)
        self.assertIn("Heterorrelato — Maria Silva", text)
        self.assertIn("respondido por Maria Silva", text)
        self.assertIn("Maria Clara", PatientIdentityService.foreign_names_for_context("O paciente Maria Clara foi avaliado.", context))

    def test_comparison_supports_more_than_four_registered_respondents(self):
        context = self.context()
        context["validated_tests"] = [test for test in context["validated_tests"] if test["instrument_code"] != "srs2"]
        context["validated_tests"].extend(self.srs2_test("adulto_heterorrelato", f"Respondente {index}", [60 + index] * 7) for index in range(5))
        data = self.export(context)
        target = ReportExportService._document_chart_targets(data)[-1]
        with ZipFile(BytesIO(data)) as package:
            root = ET.fromstring(package.read(target))
        self.assertEqual(len(root.findall(".//c:ser", ReportExportService.CHART_NS)), 5)

    def test_saved_bdefs_and_iphexa_results_are_presented_without_model_scores(self):
        context = self.context()
        context["validated_tests"].extend([
            {"instrument_code": "bdefs", "report_results": [
                {"scale": "Organização", "raw_score": 18, "percentil": 40},
                {"scale": "Total", "raw_score": 95, "percentil": 60},
            ]},
            {"instrument_code": "iphexa", "structured_results": {"dimensions": {
                "Honestidade": {"score": 0, "percentile": 10},
                "Emotividade": {"score": 3, "percentile": 65},
                "Extroversão": {"score": 4, "percentile": 80},
            }}},
        ])
        document = Document(BytesIO(self.export(context)))
        text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        self.assertIn("Organização", text)
        self.assertIn("Honestidade", text)
        row = next(row for table in document.tables for row in table.rows if row.cells[0].text == "Honestidade")
        self.assertEqual(row.cells[1].text, "0")
        self.assertEqual(row.cells[2].text, "10")
        captions = [paragraph.text for paragraph in document.paragraphs]
        self.assertTrue(any("BDEFS - Resultado total" in caption for caption in captions))
        self.assertTrue(any("IPHEXA - Radar dos resultados" in caption for caption in captions))

    def test_other_primary_models_do_not_use_the_wais3_builder(self):
        for primary in ("wisc4", "wasi"):
            with self.subTest(primary=primary), patch("apps.reports.builders.wais3_docx_builder.WAIS3DocxBuilder.build") as build:
                self.export(self.context(primary, 12))
                build.assert_not_called()

    def test_wais3_generated_domains_follow_clinical_outline_and_correct_scaled_labels(self):
        context = self.context()
        context["validated_tests"][0]["structured_results"]["subtestes"] = {key: {"escore_ponderado": score} for key, score in (("semelhancas", 16), ("vocabulario", 8), ("compreensao", 7))}
        text = WAIS3StandardizationService.build("linguagem", context)
        self.assertIn("PP=8, faixa Média", text)
        self.assertIn("PP=7, faixa Média Inferior", text)
        self.assertIn("Em análise clínica", text)
        self.assertIn("heterogêneo", text)
        self.assertGreaterEqual(text.count("\n\n"), 3)
        rules = SectionContextService._wais3_model_rules("conclusao", context)
        self.assertTrue(rules["preserve_all_respondents"])
        self.assertTrue(rules["do_not_copy_example_diagnoses"])
